#!/usr/bin/env python3
"""
check_deck.py - lint a built deck for spacing, sizing and layout slop.

Run it on every deck, every time, before you look at the slides:

    uv run --with python-pptx python scripts/check_deck.py out/deck.pptx

ERRORs must be fixed. WARNings need a reason to keep. The checks encode the
mistakes these decks actually make: text too small to read on a projector, panels
trailing dead space, blank bullets, content colliding with the footer.

Exit code is 1 if any ERROR was found, else 0.
"""
from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

SLIDE_W, SLIDE_H = 10.0, 5.625
FOOTER_Y = 5.10          # footer + slide number band
TITLE_BAND = 1.35        # titles live above this
DASH_Y = 1.39            # the layout's dashed divider under the title
CONTENT_TOP = 1.50       # content starting above this crowds the divider
MIN_PT_LIGHT = 12.0      # minimum readable size on the white surface
MIN_PT_MOCK = 8.0        # inside a dark mock window, UI chrome may be smaller
MAX_DEAD_SPACE = 1.30    # inches of empty canvas below the last content
MIN_GAP = 0.04           # shapes closer than this are touching

DARK_FILLS = {"303138", "25262C", "3C3D46"}
ACCENTS = {"E91D63", "ED1763"}


def inches(v):
    return Emu(v).inches if v is not None else 0.0


class Report:
    def __init__(self):
        self.items = defaultdict(list)

    def add(self, slide_no, level, code, msg):
        self.items[slide_no].append((level, code, msg))

    @property
    def errors(self):
        return sum(1 for v in self.items.values() for lv, _, _ in v if lv == "ERROR")

    @property
    def warnings(self):
        return sum(1 for v in self.items.values() for lv, _, _ in v if lv == "WARN")

    def render(self, total_slides):
        if not self.items:
            print(f"clean - {total_slides} slides, nothing to fix")
            return
        for no in sorted(self.items):
            print(f"\nslide {no}")
            for level, code, msg in sorted(self.items[no]):
                mark = "!" if level == "ERROR" else "-"
                print(f"  {mark} [{code}] {msg}")
        print(f"\n{self.errors} error(s), {self.warnings} warning(s), "
              f"{total_slides} slides")


def shape_box(sh, layout=None):
    """Effective geometry, resolving inherited placeholder position.

    A cloned placeholder often has no xfrm of its own and reports None for
    left/top/width/height. Treating that as 0 would make every geometry check
    silently pass, so fall back to the layout's placeholder.
    """
    left, top, width, height = sh.left, sh.top, sh.width, sh.height
    if None in (left, top, width, height) and sh.is_placeholder and layout is not None:
        idx = sh.placeholder_format.idx
        for lph in layout.placeholders:
            if lph.placeholder_format.idx == idx:
                left = left if left is not None else lph.left
                top = top if top is not None else lph.top
                width = width if width is not None else lph.width
                height = height if height is not None else lph.height
                break
    return (inches(left), inches(top), inches(width), inches(height))


def is_in_mock(sh, dark_regions):
    x, y, w, h = shape_box(sh)
    cx, cy = x + w / 2, y + h / 2
    return any(rx <= cx <= rx + rw and ry <= cy <= ry + rh
               for rx, ry, rw, rh in dark_regions)


def text_of(sh):
    if not sh.has_text_frame:
        return ""
    return sh.text_frame.text


def est_text_height(sh, layout=None):
    """Very rough: how tall the text wants to be, given the box width."""
    if not sh.has_text_frame:
        return 0.0
    _, _, w, _ = shape_box(sh, layout)
    if w <= 0:
        return 0.0
    total = 0.0
    for para in sh.text_frame.paragraphs:
        txt = "".join(r.text for r in para.runs)
        size = 18.0                      # inherited body size from the master
        for r in para.runs:
            if r.font.size:
                size = r.font.size.pt
                break
        per_line = max(1, int(w * 72 / (size * 0.55)))
        lines = max(1, -(-len(txt) // per_line)) if txt else 1
        total += lines * size * 1.25 / 72
    return total


def check(path: Path) -> int:
    prs = Presentation(str(path))
    rep = Report()
    slides = list(prs.slides)

    for no, slide in enumerate(slides, 1):
        shapes = [sh for sh in slide.shapes]

        # regions of dark mock chrome, so we can relax rules inside them
        dark_regions = []
        for sh in shapes:
            try:
                if sh.fill.type == 1 and str(sh.fill.fore_color.rgb) in DARK_FILLS:
                    dark_regions.append(shape_box(sh))
            except Exception:
                pass

        content_bottom = 0.0
        content_top = SLIDE_H
        has_title = False

        for sh in shapes:
            x, y, w, h = shape_box(sh)
            txt = text_of(sh)
            is_slide_no = (sh.is_placeholder
                           and sh.placeholder_format.idx == 12)
            if sh.is_placeholder and "TITLE" in str(sh.placeholder_format.type):
                has_title = True
            elif sh.has_text_frame and txt.strip() and y < TITLE_BAND:
                # a composed heading (e.g. the full-page closer) is still a title
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        if r.font.size and r.font.size.pt >= 24:
                            has_title = True

            # --- bounds -------------------------------------------------
            if x < -0.01 or y < -0.01 or x + w > SLIDE_W + 0.01 or y + h > SLIDE_H + 0.01:
                rep.add(no, "ERROR", "off-canvas",
                        f"{describe(sh)} runs outside the slide "
                        f"({x:.2f},{y:.2f} {w:.2f}x{h:.2f})")

            if not is_slide_no and y + h > FOOTER_Y + 0.02 and txt.strip():
                rep.add(no, "ERROR", "footer-collision",
                        f"{describe(sh)} reaches y={y + h:.2f}\", crossing the "
                        f"footer band at {FOOTER_Y}\"")

            # --- type size ----------------------------------------------
            in_mock = is_in_mock(sh, dark_regions)
            floor = MIN_PT_MOCK if in_mock else MIN_PT_LIGHT
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        if r.font.size and r.font.size.pt < floor and r.text.strip():
                            where = "inside the mock" if in_mock else "on the white surface"
                            rep.add(no, "ERROR", "text-too-small",
                                    f"{r.font.size.pt:g}pt {where} "
                                    f"(min {floor:g}pt): {short(r.text)}")

            # --- empty paragraphs (ragged bullet gaps) -------------------
            if sh.is_placeholder and sh.has_text_frame:
                paras = sh.text_frame.paragraphs
                for i, para in enumerate(paras):
                    body = "".join(r.text for r in para.runs)
                    if not body.strip() and 0 < i < len(paras) - 1:
                        rep.add(no, "WARN", "blank-bullet",
                                "empty paragraph inside a placeholder renders as a "
                                "blank bullet - pass a heading argument instead")
                        break

            # --- a single point must not wear a bullet -------------------
            if (sh.is_placeholder and sh.has_text_frame
                    and sh.placeholder_format.idx in (1, 2)
                    and "TWO_COLUMNS" in slide.slide_layout.name):
                paras = [p for p in sh.text_frame.paragraphs
                         if "".join(r.text for r in p.runs).strip()]
                if len(paras) == 1 and not _has_bullet_off(paras[0]):
                    rep.add(no, "WARN", "lone-bullet",
                            f"single point rendered as a bullet: "
                            f"{short(paras[0].text)} - suppress the glyph")

            # --- leftover hanging indent ---------------------------------
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    if not _has_bullet_off(para):
                        continue
                    pPr = para._p.find(f"{NS}pPr")
                    marL = int(pPr.get("marL") or 0)
                    ind = int(pPr.get("indent") or 0)
                    if marL or ind:
                        rep.add(no, "WARN", "hanging-indent",
                                "bullet glyph removed but its indent remains - the "
                                "copy will render ragged; zero marL and indent")
                        break

            # --- text overflowing its box --------------------------------
            # Body placeholders are checked too: s.beside() and s.body() write into
            # them, so skipping placeholders left the most common copy block on a
            # slide unmeasured. Titles are exempt (bottom-anchored, generous box);
            # so is the slide-number placeholder.
            body_ph = (sh.is_placeholder
                       and sh.placeholder_format.idx in (1, 2))
            if txt.strip() and sh.has_text_frame and (not sh.is_placeholder or body_ph):
                want = est_text_height(sh, slide.slide_layout if body_ph else None)
                if want > h + 0.08:
                    rep.add(no, "WARN", "text-overflow",
                            f"{describe(sh)} needs about {want:.2f}\" but the box "
                            f"is {h:.2f}\": {short(txt)}")

            if txt.strip() and not is_slide_no:
                content_bottom = max(content_bottom, y + h)
                content_top = min(content_top, y)
            elif not txt.strip() and not is_slide_no:
                content_bottom = max(content_bottom, y + h)
                content_top = min(content_top, y)

        # --- dead space -------------------------------------------------
        # Dividers and openers are meant to be sparse; everything else is not.
        sparse_by_design = slide.slide_layout.name in (
            "SECTION_HEADER", "TITLE", "MAIN_POINT", "BIG_NUMBER")
        if content_bottom > 0 and not sparse_by_design:
            dead = FOOTER_Y - content_bottom
            if dead > MAX_DEAD_SPACE:
                rep.add(no, "WARN", "dead-space",
                        f"{dead:.2f}\" of empty canvas below the content - add "
                        f"substance, enlarge the visual, or call .fit()")

        # --- title ------------------------------------------------------
        if not has_title and len(shapes) > 1:
            rep.add(no, "WARN", "no-title", "slide has no title placeholder")


        # --- the layout's dashed divider ---------------------------------
        if _layout_has_dash(slide.slide_layout):
            for sh in shapes:
                x, y, w, h = shape_box(sh, slide.slide_layout)
                if sh.is_placeholder or h <= 0:
                    continue
                # content that starts just under the dash reads as crowding it
                if DASH_Y < y < CONTENT_TOP:
                    rep.add(no, "WARN", "crowds-divider",
                            f'{describe(sh)} starts at y={y:.2f}", too close to the '
                            f'dashed divider at {DASH_Y}" - start content at '
                            f'{CONTENT_TOP}" or below')
                    break
            for sh in shapes:
                x, y, w, h = shape_box(sh, slide.slide_layout)
                if h > 0.10 or w < 0.25 or y > 2.2:
                    continue
                if sh.has_text_frame and text_of(sh).strip():
                    continue
                try:
                    if sh.fill.type == 1 and str(sh.fill.fore_color.rgb) in ACCENTS:
                        rep.add(no, "WARN", "double-rule",
                                "this layout already draws a dashed divider under "
                                "the title - drop the accent rule, or build on BLANK")
                        break
                except Exception:
                    pass

        # --- overlapping text -------------------------------------------
        text_shapes = [sh for sh in shapes
                       if text_of(sh).strip()
                       and not (sh.is_placeholder and sh.placeholder_format.idx == 12)]
        for i, a in enumerate(text_shapes):
            for b in text_shapes[i + 1:]:
                ov = overlap(ink_box(a, slide.slide_layout),
                             ink_box(b, slide.slide_layout))
                if ov > 0.12:
                    rep.add(no, "WARN", "overlap",
                            f"{describe(a)} and {describe(b)} overlap by "
                            f"{ov:.2f} sq in")

        # --- dark tables ------------------------------------------------
        for sh in shapes:
            if getattr(sh, "has_table", False) and sh.has_table:
                if is_in_mock(sh, dark_regions):
                    continue        # mocked product output is dark on purpose
                for r in sh.table.rows:
                    for c in r.cells:
                        try:
                            if str(c.fill.fore_color.rgb) in DARK_FILLS:
                                rep.add(no, "WARN", "dark-table",
                                        "table uses the dark mock palette - tables "
                                        "are slide content and belong on white")
                                break
                        except Exception:
                            pass
                    else:
                        continue
                    break


    rep.render(len(slides))
    return 1 if rep.errors else 0


def _layout_has_dash(layout) -> bool:
    """True when the layout draws its own dashed divider under the title."""
    return 'prstDash val="lgDash"' in layout._element.xml


NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _has_bullet_off(para) -> bool:
    """True when the paragraph explicitly suppresses its inherited bullet."""
    pPr = para._p.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}pPr")
    if pPr is None:
        return False
    return pPr.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}buNone") is not None


def ink_box(sh, layout=None):
    """The box the text actually paints, not the box it was given.

    A short left-aligned title sits in a full-width placeholder; comparing raw
    boxes reports collisions that no viewer can see. Narrow the box to the
    estimated text extent when the text is too short to wrap.
    """
    x, y, w, h = shape_box(sh, layout)
    if not sh.has_text_frame or w <= 0:
        return (x, y, w, h)
    widest = 0.0
    for para in sh.text_frame.paragraphs:
        txt = "".join(r.text for r in para.runs)
        if not txt.strip():
            continue
        size = 18.0                      # inherited body size from the master
        for r in para.runs:
            if r.font.size:
                size = r.font.size.pt
                break
        if sh.is_placeholder and "TITLE" in str(sh.placeholder_format.type):
            size = 30.0                      # inherited from the master
        widest = max(widest, len(txt) * size * 0.5 / 72)
    if widest and widest < w:
        return (x, y, widest, h)
    return (x, y, w, h)


def overlap(a, b):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    dx = min(ax + aw, bx + bw) - max(ax, bx)
    dy = min(ay + ah, by + bh) - max(ay, by)
    return dx * dy if dx > MIN_GAP and dy > MIN_GAP else 0.0


def describe(sh):
    t = text_of(sh).strip().replace("\n", " ")
    if t:
        return f'"{short(t)}"'
    return sh.shape_type and str(sh.shape_type).split()[0].lower() or "shape"


def short(t, n=38):
    t = " ".join(str(t).split())
    return t if len(t) <= n else t[: n - 1] + "…"


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("usage: check_deck.py <deck.pptx> [more.pptx ...]")
    rc = 0
    for arg in sys.argv[1:]:
        p = Path(arg)
        print(f"== {p.name}")
        rc |= check(p)
    sys.exit(rc)
