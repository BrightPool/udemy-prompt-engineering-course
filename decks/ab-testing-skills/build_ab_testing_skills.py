#!/usr/bin/env python3
"""
"A/B Testing - How to Systematically Improve Your Skills" - the principles half
of a new lecture. Closes the Skills & Plugins section.

Run:
    uv run --with python-pptx python decks/ab-testing-skills/build_ab_testing_skills.py \
        decks/ab-testing-skills/ab-testing-skills.pptx
    uv run --with python-pptx python .claude/skills/principles-deck/scripts/check_deck.py \
        decks/ab-testing-skills/ab-testing-skills.pptx

Boundaries. 'What are Skills?' has already taught what a skill is and that its
description decides when it fires; the two 'Creating a skill' lectures build one.
Mike's GSheets lecture is the mechanics of running a prompt test in a spreadsheet.
This deck is the *method* of improving a skill, and it teaches three things:

  1. Change one thing. Change two and you learn nothing, however much better
     the output looks.
  2. Write the criteria before you see the outputs, or you will rationalise
     whichever version you already preferred.
  3. Judge across samples. These systems vary run to run, so one good answer is
     not evidence - which is why ten.

Running example throughout: a `meeting-notes` skill, and one line added to it.
Deliberately no eval framework, no code, no tooling - a tally is the method.
"""
import sys
from pathlib import Path

SKILL = Path(__file__).resolve().parents[2] / ".claude/skills/principles-deck/scripts"
sys.path.insert(0, str(SKILL))
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from deckkit import (Deck, INK, ACCENT, BODY_FONT, title_case, _rgb,  # noqa: E402
                     SOFT_INK, RULE, NODE_ACC, NODE_BG, SZ_CAPTION)

ASSETS = Path(__file__).resolve().parent / "assets"


def icon(s, name, x, y, size, color="424242"):
    """A bare Lucide icon (ISC licence), pre-rendered in assets/."""
    return s.raw.shapes.add_picture(str(ASSETS / f"icon-{name}-{color}.png"),
                                    Inches(x), Inches(y), Inches(size), Inches(size))


def badge(s, name, x, y, d, fill, color, line=None):
    """A filled circle with an icon centred in it."""
    o = s.raw.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                               Inches(d), Inches(d))
    o.fill.solid()
    o.fill.fore_color.rgb = _rgb(fill)
    if line:
        o.line.color.rgb = _rgb(line)
        o.line.width = Pt(1.0)
    else:
        o.line.fill.background()
    o.shadow.inherit = False
    g = d * 0.56
    icon(s, name, x + (d - g) / 2, y + (d - g) / 2, g, color)


def dot(s, x, y, d, fill, line=None):
    o = s.raw.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                               Inches(d), Inches(d))
    o.fill.solid()
    o.fill.fore_color.rgb = _rgb(fill)
    if line:
        o.line.color.rgb = _rgb(line)
        o.line.width = Pt(1.0)
    else:
        o.line.fill.background()
    o.shadow.inherit = False


class grouped:
    """`with grouped(s, "name"):` groups every shape drawn inside the block."""

    def __init__(self, s, name):
        self.s, self.name = s, name

    def __enter__(self):
        self.mark = len(self.s.raw.shapes)

    def __exit__(self, *exc):
        new = list(self.s.raw.shapes)[self.mark:]
        if new and not exc[0]:
            self.s.raw.shapes.add_group_shape(new).name = self.name


def tc_table(rows):
    """Header row and first-column row labels are labels, so Title Case them."""
    return [[title_case(c) if (r == 0 or j == 0) and c else c
             for j, c in enumerate(row)] for r, row in enumerate(rows)]


out = Path(sys.argv[1] if len(sys.argv) > 1
           else Path(__file__).resolve().parent / "ab-testing-skills.pptx")
d = Deck()

# ---------------------------------------------------------------- open
d.title("A/B Testing Your Skills",
        "How to tell whether a change to a skill made it better, or only "
        "made it different.")

# ============================================================== WHAT IS IT?
s = d.what("the same job, two versions, one line apart")
m = s.mock(0.40, 1.52, 4.60, 3.58, generic=True)
m.label("Two versions, one job")
m.user("Turn this transcript into notes.")
m.divider()
m.label("Version A")
m.line("- Chased the vendor about pricing")
m.line("- Someone to redo the deck")
m.gap(0.08)
m.label("Version B - one line added")
m.line("Decisions: ship on the 14th; drop tier 3.")
m.line("Actions: Priya - redo the deck by Friday.", highlight=True)
m.fit()
s.beside(m, "Two versions, one line apart. Judged over a set of cases, against "
            "criteria written first.")

# Three runs of the same thing, drawn as counts, so the variation is visible.
s = d.content("Why one good answer proves nothing")
with grouped(s, "Same input strip"):
    s.rect(0.34, 1.52, 9.32, 0.44, fill=NODE_BG, rounded=True, radius=0.10,
           line=RULE, line_w=0.75)
    icon(s, "dices", 0.50, 1.60, 0.28, "424242")
    s.text(0.90, 1.52, 8.60, 0.44,
           "Same transcript. Same skill. Nothing changed between runs.",
           size=SZ_CAPTION, color=INK, font=BODY_FONT, anchor="m")
runs = [("Run 1", (3, 5, 0), "All actions owned", False),
        ("Run 2", (3, 7, 1), "One action unowned", True),
        ("Run 3", (2, 5, 0), "Merged two decisions", True)]
cw, gap, cy, ch = 2.96, 0.22, 2.12, 2.30
for i, (name, counts, note, off) in enumerate(runs):
    x = 0.34 + i * (cw + gap)
    with grouped(s, name):
        s.rect(x, cy, cw, ch, fill="FFFFFF", rounded=True, radius=0.06,
               line=ACCENT if off else RULE, line_w=1.0 if off else 0.75)
        s.text(x + 0.18, cy + 0.14, cw - 0.36, 0.34, name, size=14, color=INK,
               bold=True, font=BODY_FONT, anchor="m")
        for r, (label, n) in enumerate(zip(("Decisions", "Actions", "Unowned"),
                                           counts)):
            ry = cy + 0.62 + r * 0.40
            s.text(x + 0.18, ry, 1.10, 0.26, label, size=SZ_CAPTION,
                   color=SOFT_INK, font=BODY_FONT, anchor="m")
            for k in range(n):
                dot(s, x + 1.30 + k * 0.20, ry + 0.06, 0.14,
                    ACCENT if label == "Unowned" else INK)
            if n == 0:
                s.text(x + 1.30, ry, 0.60, 0.26, "none", size=SZ_CAPTION,
                       color=SOFT_INK, font=BODY_FONT, anchor="m")
        s.text(x + 0.18, cy + ch - 0.44, cw - 0.36, 0.30, note, size=SZ_CAPTION,
               color=ACCENT if off else SOFT_INK, bold=off, font=BODY_FONT,
               anchor="m")
s.caption(0.34, 4.60, 9.32,
          "One run is a draw, not a result. Judge how often each version gets it right.")

# Ten cases, drawn as ten tokens: a useless ten against a useful ten.
s = d.content("Choosing the ten cases")
TW, TG = 0.86, 0.074


def case_row(y, head, tokens, bad):
    with grouped(s, head):
        s.text(0.34, y, 9.32, 0.28, title_case(head), size=13,
               color=SOFT_INK if bad else ACCENT, bold=True, font=BODY_FONT)
        for k, (label, style) in enumerate(tokens):
            x = 0.34 + k * (TW + TG)
            f, ln, tc = {"easy": ("FFFFFF", RULE, SOFT_INK),
                         "real": (NODE_BG, RULE, INK),
                         "hard": (NODE_ACC, ACCENT, ACCENT),
                         "fail": (ACCENT, ACCENT, "FFFFFF")}[style]
            s.rect(x, y + 0.36, TW, 0.46, fill=f, rounded=True, radius=0.10,
                   line=ln, line_w=0.75)
            s.text(x, y + 0.36, TW, 0.46, label, size=SZ_CAPTION, color=tc,
                   bold=(style != "easy"), font=BODY_FONT, align="c", anchor="m")


case_row(1.52, "Ten that teach you nothing", [("Easy", "easy")] * 10, True)
case_row(2.52, "Ten worth running",
         [("Real", "real")] * 6 + [("Hard", "hard")] * 3 + [("Failed", "fail")],
         False)
with grouped(s, "Legend"):
    lx = 0.34
    for label, f, ln in (("Real work you did", NODE_BG, RULE),
                         ("Where it usually slips", NODE_ACC, ACCENT),
                         ("The one that went wrong", ACCENT, ACCENT)):
        s.rect(lx, 3.50, 0.20, 0.20, fill=f, rounded=True, radius=0.04,
               line=ln, line_w=0.75)
        s.text(lx + 0.28, 3.44, 2.90, 0.32, label, size=SZ_CAPTION, color=INK,
               font=BODY_FONT, anchor="m")
        lx += 3.10
with grouped(s, "Frozen"):
    s.rect(0.34, 4.02, 9.32, 0.50, fill="FFFFFF", rounded=True, radius=0.08,
           line=RULE, line_w=0.75)
    badge(s, "lock", 0.46, 4.08, 0.38, NODE_ACC, "E91D63")
    s.text(1.00, 4.02, 8.50, 0.50,
           "Pick the ten before the change, then freeze the list.",
           size=SZ_CAPTION, color=INK, font=BODY_FONT, anchor="m")

# ============================================================== HOW IT WORKS
s = d.how("hold everything else still")
steps = [("pencil-line", "Write Criteria", "what better means"),
         ("lock", "Freeze It", "that is the control"),
         ("copy", "Change a Line", "that is the variant"),
         ("play", "Run Both", "on the same ten"),
         ("chart-column", "Score and Count", "against step 1")]
col = 9.32 / len(steps)
BD = 0.70
s.rect(0.34 + col / 2, 1.70 + BD / 2, col * (len(steps) - 1), 0.02, fill=RULE)
for i, (ico, label, sub) in enumerate(steps):
    cx = 0.34 + i * col + col / 2
    last = i == len(steps) - 1
    with grouped(s, label):
        badge(s, ico, cx - BD / 2, 1.70, BD, ACCENT if last else "FFFFFF",
              "FFFFFF" if last else "424242", line=None if last else RULE)
        # The step number rides on the badge, so the label fits on one line.
        s.disc(cx + BD / 2 - 0.20, 1.62, 0.30, str(i + 1), size=12)
        s.text(cx - col / 2 + 0.04, 2.54, col - 0.08, 0.30, label,
               size=12, color=ACCENT if last else INK, bold=True,
               font=BODY_FONT, align="c")
        s.text(cx - col / 2 + 0.04, 2.86, col - 0.08, 0.30, title_case(sub),
               size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT, align="c")
s.caption(0.34, 3.66, 9.32,
          "Step 1 must come before step 4, or the tally is just you agreeing "
          "with yourself.")

s = d.content("The one line that moved")
s.text(0.34, 1.52, 4.46, 0.26, title_case("Control"), size=14, color=INK, bold=True,
       font=BODY_FONT)
s.text(5.20, 1.52, 4.46, 0.26, title_case("Variant"), size=14, color=ACCENT, bold=True,
       font=BODY_FONT)
fixed = [("The request", "same wording"),
         ("The ten cases", "same ten"),
         ("The model", "same one")]
s.layers(0.34, 1.90, 4.46, fixed + [("The skill", "as it is today")], h=0.50)
s.layers(5.20, 1.90, 4.46, fixed + [("The skill", "one line added", True)],
         h=0.50)
s.caption(0.34, 4.28, 9.32,
          "Change two lines and a better result cannot tell you which one did it.")

s = d.content("Criteria, written before you look")
r = s.message_row(0.34, 1.52, 9.31, 1.41)
r.label("Written before the run")
r.line("Every decision in the transcript appears, and is labelled a decision.")
r.line("Every action item names an owner.")
r.line("Nothing appears that nobody said.")
r.line("Under 200 words.")
r = s.message_row(0.34, 3.10, 9.31, 0.72)
r.label("Written after the run")
r.line("Version B just reads better.", highlight=True)
with grouped(s, "Blind scoring"):
    badge(s, "shuffle", 0.34, 4.02, 0.40, NODE_ACC, "E91D63")
    s.text(0.86, 4.02, 8.80, 0.40,
           "Only criteria written first can be scored blind. Strip the labels "
           "and shuffle.", size=SZ_CAPTION, color=INK, font=BODY_FONT, anchor="m")

# The tally as a bar chart: control against variant, out of ten.
s = d.content("Ten runs, one tally")
tally = [("Every decision listed", 6, 9, "Win"),
         ("Every action owned", 4, 9, "Win"),
         ("Nothing invented", 10, 8, "Weigh"),
         ("Under 200 words", 9, 6, "Loss")]
BX, BW = 3.10, 4.90             # bar origin and the width of 10 out of 10
with grouped(s, "Legend"):
    for k, (label, f) in enumerate((("Control", "C9C9CF"), ("Variant", ACCENT))):
        s.rect(BX + k * 1.50, 1.56, 0.20, 0.20, fill=f)
        s.text(BX + k * 1.50 + 0.28, 1.50, 1.10, 0.32, label, size=SZ_CAPTION,
               color=INK, font=BODY_FONT, anchor="m")
for i, (label, c, v, verdict) in enumerate(tally):
    y = 2.00 + i * 0.66
    with grouped(s, label):
        s.text(0.34, y, 2.70, 0.52, title_case(label), size=SZ_CAPTION, color=INK,
               bold=True, font=BODY_FONT, anchor="m")
        for k, (n, f) in enumerate(((c, "C9C9CF"), (v, ACCENT))):
            by = y + 0.04 + k * 0.24
            s.rect(BX, by, BW * n / 10, 0.20, fill=f)
            s.text(BX + BW * n / 10 + 0.08, by - 0.04, 0.50, 0.28, str(n),
                   size=SZ_CAPTION, color=INK, font=BODY_FONT, anchor="m")
        vf, vt = {"Win": (ACCENT, "FFFFFF"), "Loss": ("FFFFFF", ACCENT),
                  "Weigh": (NODE_BG, INK)}[verdict]
        s.rect(8.66, y + 0.08, 1.00, 0.36, fill=vf, rounded=True, radius=0.18,
               line=ACCENT if verdict != "Weigh" else RULE, line_w=0.75)
        s.text(8.66, y + 0.08, 1.00, 0.36, verdict, size=SZ_CAPTION, color=vt,
               bold=True, font=BODY_FONT, align="c", anchor="m")
s.caption(0.34, 4.66, 9.32,
          "The loss is the useful part: the variant started padding. A one-case "
          "gap is noise.")

s = d.content("Three results, three decisions")
gap, w = 0.30, (9.32 - 0.30 * 2) / 3
for i, (ico, label, sub) in enumerate([
        ("trophy", "It clearly won", "adopt it, and accept what it cost"),
        ("circle-x", "It clearly lost", "you avoided a habit that does nothing"),
        ("scale", "Too close to call", "keep the control, change something bigger")]):
    x = 0.34 + i * (w + gap)
    win = i == 0
    with grouped(s, label):
        s.rect(x, 1.56, w, 2.10, fill=NODE_ACC if win else "FFFFFF", rounded=True,
               radius=0.06, line=ACCENT if win else RULE,
               line_w=1.0 if win else 0.75)
        badge(s, ico, x + w / 2 - 0.34, 1.76, 0.68, ACCENT if win else NODE_BG,
              "FFFFFF" if win else "424242", line=None if win else RULE)
        s.text(x + 0.14, 2.60, w - 0.28, 0.32, title_case(label), size=14,
               color=ACCENT if win else INK, bold=True, font=BODY_FONT, align="c")
        s.text(x + 0.18, 2.96, w - 0.36, 0.56, title_case(sub), size=SZ_CAPTION,
               color=SOFT_INK, font=BODY_FONT, align="c", spacing=1.15)
s.caption(0.34, 3.90, 9.32,
          "All three are results. A change that did nothing is worth knowing too.")

# ============================================================== WHY IT MATTERS
s = d.why("the skill gets better on purpose")
s.flow(0.34, 1.52, 9.32,
       [("Change one line", "in the skill"),
        ("Run the ten", "both versions"),
        ("Score and count", "against your criteria"),
        ("Keep the winner", "as the new control")],
       h=1.10, accent_last=True,
       loop="then move the next single line")
s.caption(0.34, 3.60, 9.32,
          "Run the loop a few times, one line per pass. You end up knowing why "
          "every line is there.")

s = d.content("What the tally does not tell you")
s.table(0.34, 1.52, 9.32, 2.00, tc_table([
    ["Your result", "What it supports", "What it does not"],
    ["Coverage", "These ten cases", "Cases you did not think of"],
    ["Shelf life", "The model as it is today", "The model after its next update"],
    ["Strength", "A gap you can see across ten", "A one- or two-case difference"],
]), col_widths=[2.0, 3.6, 3.7])
s.caption(0.34, 3.68, 9.32,
          "A won test is evidence, not proof: this version, your cases, this week.")

# When to bother: two panels, each with its own icon.
s = d.content("When it is worth ten runs")
panels = [("flask-conical", "Worth testing", True,
           ["A skill the team runs every week",
            "Two wordings, no way to choose",
            "A change that slows it down",
            "A rule someone insists on adding"]),
          ("pencil", "Just change it", False,
           ["A plain mistake in the instructions",
            "A skill you will use once",
            "A change you can check in one read",
            "A result you would ignore either way"])]
pw, pg = (9.32 - 0.30) / 2, 0.30
for i, (ico, head, accent, items) in enumerate(panels):
    x = 0.34 + i * (pw + pg)
    with grouped(s, head):
        s.rect(x, 1.56, pw, 3.10, fill="FFFFFF", rounded=True, radius=0.05,
               line=ACCENT if accent else RULE, line_w=1.0 if accent else 0.75)
        badge(s, ico, x + 0.20, 1.74, 0.50, ACCENT if accent else NODE_BG,
              "FFFFFF" if accent else "424242", line=None if accent else RULE)
        s.text(x + 0.84, 1.74, pw - 1.0, 0.50, title_case(head), size=15,
               color=ACCENT if accent else INK, bold=True, font=BODY_FONT,
               anchor="m")
        for k, item in enumerate(items):
            iy = 2.52 + k * 0.50
            s.rect(x + 0.24, iy + 0.12, 0.08, 0.08,
                   fill=ACCENT if accent else SOFT_INK)
            s.text(x + 0.44, iy, pw - 0.64, 0.34, item, size=SZ_CAPTION,
                   color=INK, font=BODY_FONT, anchor="m")

# ---------------------------------------------------------------- handoff
d.handoff(
    [("Write criteria", "before the change"),
     ("Change one line", "and nothing else"),
     ("Run ten cases", "both versions"),
     ("Tally and decide", "adopt or discard")],
    lead="That is the method. Now we run one, live, on a real skill.",
)

path = d.save(out)
print(f"wrote {path}  ({len(d.prs.slides._sldIdLst)} slides)")
