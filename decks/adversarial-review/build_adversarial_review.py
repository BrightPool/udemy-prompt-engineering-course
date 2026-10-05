#!/usr/bin/env python3
"""Build the principles deck for Adversarial Review of Outputs."""
import sys
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / ".claude" / "skills" / "principles-deck" / "scripts"))
from deckkit import (Deck, ACCENT, INK, SOFT_INK, RULE, NODE_ACC,  # noqa: E402
                     NODE_BG, BODY_FONT, SZ_CAPTION, title_case, _rgb)

def tc_table(rows):
    """Header row and first-column row labels are labels, so Title Case them."""
    return [[title_case(c) if (r == 0 or j == 0) and c else c
             for j, c in enumerate(row)] for r, row in enumerate(rows)]


ASSETS = Path(__file__).resolve().parent / "assets"   # Lucide icons, ISC licence


def badge(s, x, y, d, icon, fill=NODE_ACC, line=ACCENT, glyph=0.55):
    """A circle with an icon centred in it."""
    c = s.raw.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                               Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = _rgb(fill)
    c.line.color.rgb = _rgb(line)
    c.line.width = Pt(1.0)
    c.shadow.inherit = False
    g = d * glyph
    s.raw.shapes.add_picture(str(ASSETS / icon), Inches(x + (d - g) / 2),
                             Inches(y + (d - g) / 2), Inches(g), Inches(g))


def icon(s, name, x, y, size):
    s.raw.shapes.add_picture(str(ASSETS / name), Inches(x), Inches(y),
                             Inches(size), Inches(size))


def group(s, start, name):
    """Group every shape added since `start`, so it moves as one unit."""
    shapes = list(s.raw.shapes)[start:]
    if len(shapes) > 1:
        s.raw.shapes.add_group_shape(shapes).name = name


CODEX_ICON = ROOT / "decks" / "codex-goal" / "assets" / "icon-codex-light.png"

out = Path(sys.argv[1] if len(sys.argv) > 1 else
           Path(__file__).resolve().parent / "adversarial-review.pptx")
d = Deck()

d.title(
    "Adversarial Review of Outputs",
    "Have a second AI check the work against clear criteria, so you only review "
    "what has already passed.",
)

# 1. What: the first output is a draft, however confident it sounds.
s = d.what("fluent is not the same as correct")
m = s.mock(0.40, 1.56, 4.60, 3.40, generic=True)
m.user("Summarise Q3 churn from the attached sheet.")
with m.bubble():
    m.assistant("Q3 churn fell to 3.1%.", highlight=True)
    m.line("Every region improved.", highlight=True)
    m.line("The annual plan drove the drop.")
m.fit()
fx = 5.40
s.raw.shapes.add_picture(str(ASSETS / "icon-sheet-424242.png"), Inches(fx),
                         Inches(1.60), Inches(0.30), Inches(0.30))
s.text(fx + 0.42, 1.58, 3.80, 0.34, title_case("Checked against the sheet"),
       size=14, color=INK, bold=True, font=BODY_FONT, anchor="m")
s.rect(fx, 2.02, 4.26, 0.012, fill=INK)
checks = [("icon-circle-x-E91D63.png", "\"Fell to 3.1%\"", "Sheet says 4.3%"),
          ("icon-circle-x-E91D63.png", "\"Every region improved\"", "EMEA got worse"),
          ("icon-circle-help-6B6B6B.png", "\"Annual plan drove it\"", "Not in the sheet")]
for i, (ic, claim, truth) in enumerate(checks):
    y = 2.14 + i * 0.62
    start = len(s.raw.shapes)
    icon(s, ic, fx, y + 0.06, 0.30)
    s.text(fx + 0.42, y, 3.84, 0.26, claim, size=SZ_CAPTION, color=SOFT_INK,
           font=BODY_FONT)
    s.text(fx + 0.42, y + 0.26, 3.84, 0.26, truth, size=SZ_CAPTION, color=ACCENT,
           bold=True, font=BODY_FONT)
    s.rect(fx, y + 0.56, 4.26, 0.008, fill=RULE)
    group(s, start, f"Check {i + 1}")
s.text(fx, 4.12, 4.26, 0.60, "Equally sure either way. Treat it as a draft.",
       size=14, color=INK, font=BODY_FONT, spacing=1.15)

# 2. How: maker, reviewer, criteria, loop.
s = d.how("a maker, a reviewer and a bar to clear")
stations = [("icon-pen-line-E91D63.png", "Maker Drafts", "does the task"),
            ("icon-search-check-E91D63.png", "Reviewer Checks", "against the criteria"),
            ("icon-wrench-E91D63.png", "Maker Fixes", "every failed point"),
            ("icon-user-check-FFFFFF.png", "You Review", "the final version")]
col, bd, by = 9.32 / 4, 0.84, 1.62
centres = [0.34 + col * (i + 0.5) for i in range(4)]
for i, ((ic, lab, sub), cx) in enumerate(zip(stations, centres)):
    start = len(s.raw.shapes)
    last = i == len(stations) - 1
    badge(s, cx - bd / 2, by, bd, ic, fill=ACCENT if last else NODE_ACC)
    s.text(cx - col / 2, by + bd + 0.10, col, 0.28, lab, size=14,
           color=ACCENT if last else INK, bold=True, font=BODY_FONT, align="c")
    s.text(cx - col / 2, by + bd + 0.38, col, 0.26, title_case(sub),
           size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT, align="c")
    group(s, start, lab)
    if not last:
        s.arrow(cx + bd / 2 + 0.12, by + bd / 2 - 0.06, col - bd - 0.24)
# the loop: from "Maker Fixes" back to "Reviewer Checks"
ly = by + bd + 0.80
start = len(s.raw.shapes)
s.rect(centres[2], by + bd + 0.68, 0.012, ly - (by + bd + 0.68), fill=ACCENT)
s.rect(centres[1], ly, centres[2] - centres[1], 0.012, fill=ACCENT)
s.arrow(centres[1] - 0.05, by + bd + 0.68, 0.10, ly - (by + bd + 0.68),
        direction="up")
mid = (centres[1] + centres[2]) / 2
s.text(mid - 1.80, ly + 0.06, 3.60, 0.26,
       title_case("until every criterion passes"), size=SZ_CAPTION, color=ACCENT,
       font=BODY_FONT, align="c")
group(s, start, "Review loop")
s.caption(
    0.34,
    3.72,
    9.32,
    "The reviewer gets fresh context and is told to find problems, not to agree.",
)

# 3. The key ingredient: criteria the reviewer can actually check.
s = d.content("Acceptance criteria make the review real")
m = s.mock(0.40, 1.56, 5.00, 3.44, generic=True)
m.user("Review the draft. Do not rewrite it.",
       "Mark each criterion PASS or FAIL with evidence:")
m.line("1. Every figure matches the sheet")
m.line("2. Each region has its own sentence")
m.line("3. Under 150 words, no jargon")
with m.bubble():
    m.assistant("FAIL 1: churn is 4.3%, not 3.1%.")
    m.line("FAIL 2: EMEA is missing.", highlight=True)
    m.line("PASS 3.")
m.fit()
vx = 5.84
for i, (ic, head, quote, colr) in enumerate([
    ("icon-circle-x-6B6B6B.png", "Vague", "\"Make it better\"", SOFT_INK),
    ("icon-circle-check-E91D63.png", "Checkable", "\"Every figure matches the sheet\"",
     ACCENT),
]):
    y = 1.70 + i * 1.02
    start = len(s.raw.shapes)
    icon(s, ic, vx, y, 0.36)
    s.text(vx + 0.50, y - 0.02, 3.30, 0.30, title_case(head), size=14, color=colr,
           bold=True, font=BODY_FONT)
    s.text(vx + 0.50, y + 0.30, 3.30, 0.50, quote, size=SZ_CAPTION, color=INK,
           font=BODY_FONT, spacing=1.1)
    group(s, start, head)
s.text(vx, 3.84, 3.82, 0.60, "Checkable criteria turn an opinion into a fix list.",
       size=14, color=INK, font=BODY_FONT, spacing=1.15)

# 4. Where your time goes: the two ends, not the middle.
s = d.content("You set the bar, then step away")
y, h = 1.86, 1.30
start = len(s.raw.shapes)
badge(s, 0.34 + (2.10 - 0.70) / 2, y + 0.02, 0.70, "icon-user-FFFFFF.png", fill=ACCENT)
s.text(0.34, y + 0.78, 2.10, 0.26, "You Start", size=14, color=INK, bold=True,
       font=BODY_FONT, align="c")
s.text(0.34, y + 1.04, 2.10, 0.26, title_case("goal + criteria"), size=SZ_CAPTION,
       color=SOFT_INK, font=BODY_FONT, align="c")
group(s, start, "You start")
mx, mw = 2.44 + 0.30, 9.66 - 2.10 - 0.30 - (2.44 + 0.30)
s.rect(mx, y, mw, h, fill=NODE_BG, rounded=True, radius=0.05, line=RULE,
       line_w=0.75)
icon = 0.52
s.raw.shapes.add_picture(str(CODEX_ICON), Inches(mx + 0.22),
                         Inches(y + (h - icon) / 2), Inches(icon), Inches(icon))
s.text(mx + 0.90, y + 0.22, mw - 1.10, 0.30, title_case("The loop runs without you"),
       size=13, color=INK, bold=True, font=BODY_FONT)
s.text(mx + 0.90, y + 0.58, mw - 1.10, 0.44,
       title_case("draft, review, fix, review again"), size=SZ_CAPTION, color=SOFT_INK,
       font=BODY_FONT)
start = len(s.raw.shapes)
badge(s, 9.66 - 2.10 + (2.10 - 0.70) / 2, y + 0.02, 0.70, "icon-user-check-FFFFFF.png",
      fill=ACCENT)
s.text(9.66 - 2.10, y + 0.78, 2.10, 0.26, "You Return", size=14, color=INK,
       bold=True, font=BODY_FONT, align="c")
s.text(9.66 - 2.10, y + 1.04, 2.10, 0.26, title_case("a checked result"),
       size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT, align="c")
group(s, start, "You return")
for ax in (2.44 + 0.05, mx + mw + 0.05):
    s.arrow(ax, y + h / 2 - 0.05, 0.20, 0.10)
# time axis
ty = y + h + 0.36
s.rect(0.34, ty, 9.32, 0.012, fill=SOFT_INK)
for tx, lab, al in ((0.34, "0 min", "l"), (9.66 - 4.0, "~30 min, or as long as it takes", "r")):
    s.text(tx, ty + 0.10, 4.0, 0.26, lab, wrap=False, size=SZ_CAPTION, color=SOFT_INK,
           font=BODY_FONT, align=al)
s.caption(
    0.34,
    ty + 0.62,
    9.32,
    "Your time goes into the two ends. The middle is what /goal is for.",
)

# 5. Why it matters, including how it fails.
s = d.why("it raises the floor, not the ceiling")
fails = [("icon-eye-off-E91D63.png", "Shared Blind Spots",
          "reviewer repeats the maker's mistake", "Fresh context or another model"),
         ("icon-circle-help-E91D63.png", "Vague Criteria",
          "\"looks good to me\"", "Criteria a stranger could check"),
         ("icon-stamp-E91D63.png", "Rubber-Stamping",
          "everything passes first time", "Demand evidence per criterion"),
         ("icon-infinity-E91D63.png", "Endless Loop",
          "each fix breaks something else", "Cap the rounds, then step in")]
for i, (ic, name, symptom, fix) in enumerate(fails):
    y = 1.58 + i * 0.66
    start = len(s.raw.shapes)
    badge(s, 0.34, y, 0.50, ic)
    s.text(1.00, y - 0.04, 4.10, 0.30, name, size=14, color=INK, bold=True,
           font=BODY_FONT)
    s.text(1.00, y + 0.24, 4.10, 0.26,
           symptom if symptom.startswith('"') else title_case(symptom),
           size=SZ_CAPTION, color=SOFT_INK,
           font=BODY_FONT)
    s.arrow(5.10, y + 0.20, 0.50, 0.10)
    s.text(5.80, y, 3.86, 0.50, fix, size=SZ_CAPTION, color=ACCENT, bold=True,
           font=BODY_FONT, anchor="m")
    if i < len(fails) - 1:
        s.rect(0.34, y + 0.60, 9.32, 0.008, fill=RULE)
    group(s, start, name)
s.caption(
    0.34,
    4.36,
    9.32,
    "Review catches the obvious problems. You still do the final check.",
)

d.handoff(
    [
        ("Write the criteria", "before any work starts"),
        ("Run the loop in Codex", "maker and reviewer, with /goal"),
        ("Come back and check", "the review first, then the work"),
    ],
    lead="Next, we build this loop in Codex, on top of /goal.",
)

path = d.save(out)
print(f"wrote {path}  ({len(d.prs.slides._sldIdLst)} slides)")
