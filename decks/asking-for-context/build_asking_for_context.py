#!/usr/bin/env python3
"""Build the five-slide principles deck for asking for missing context."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / ".claude" / "skills" / "principles-deck" / "scripts"))
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from deckkit import (Deck, title_case, _rgb, ACCENT, INK, SOFT_INK, RULE,  # noqa: E402
                     NODE_ACC, NODE_BG, BODY_FONT, SZ_CAPTION)

ASSETS = Path(__file__).resolve().parent / "assets"


def icon(s, name, x, y, size, color="424242"):
    """A bare Lucide icon (ISC licence), pre-rendered in assets/."""
    return s.raw.shapes.add_picture(str(ASSETS / f"icon-{name}-{color}.png"),
                                    Inches(x), Inches(y), Inches(size), Inches(size))


def badge(s, name, x, y, d, fill, color, line=None):
    """A filled circle with an icon centred in it."""
    o = s.raw.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                               Inches(d), Inches(d))
    o.fill.solid()
    o.fill.fore_color.rgb = _rgb(fill)
    if line:
        o.line.color.rgb = _rgb(line)
        o.line.width = Pt(1.0)
    else:
        o.line.fill.background()
    o.shadow.inherit = False
    g = d * 0.56
    return [o, icon(s, name, x + (d - g) / 2, y + (d - g) / 2, g, color)]


def since(s, mark):
    """Shapes added to the slide after `mark` = len(shapes) was taken."""
    return list(s.raw.shapes)[mark:]


def group(s, shapes, name):
    g = s.raw.shapes.add_group_shape(list(shapes))
    g.name = name
    return g


out = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parent / "asking-for-context.pptx")
d = Deck()

d.title(
    "Ask for Context",
    "How a short check for missing information prevents avoidable guessing.",
)

s = d.what("a context check pauses the task at the right moment")
m = s.mock(0.40, 1.52, 4.60, 3.58, generic=True)
m.user("Draft the launch email.")
m.line("Before writing, ask for missing context.", highlight=True)
with m.bubble():
    m.assistant("Who is the audience?")
    m.line("What action should they take?")
    m.line("When does the offer end?")
m.composer(pin=False)
s.beside(
    m,
    "The model names the gaps before it writes. You answer only what matters.",
)

# How it works: a request with gaps, the questions that fill them, the brief.
s = d.how("missing information becomes part of the brief")
TOP, CH = 1.62, 2.26
cols = [(0.34, 2.96), (3.64, 2.72), (6.70, 2.96)]


def card(x, w, head, ico, rows, accent=False):
    mark = len(s.raw.shapes)
    s.rect(x, TOP, w, CH, fill=NODE_ACC if accent else "FFFFFF", rounded=True,
           radius=0.06, line=ACCENT if accent else RULE, line_w=1.0 if accent else 0.75)
    badge(s, ico, x + 0.18, TOP + 0.16, 0.42, ACCENT if accent else NODE_BG,
          "FFFFFF" if accent else "424242", line=None if accent else RULE)
    s.text(x + 0.70, TOP + 0.16, w - 0.84, 0.42, title_case(head), size=14,
           color=ACCENT if accent else INK, bold=True, font=BODY_FONT, anchor="m")
    y = TOP + 0.80
    for mark_icon, colour, text in rows:
        icon(s, mark_icon, x + 0.20, y + 0.03, 0.22, colour)
        s.text(x + 0.52, y, w - 0.64, 0.28, text, size=SZ_CAPTION,
               color=ACCENT if colour == "E91D63" else INK, font=BODY_FONT,
               anchor="m", bold=(colour == "E91D63" and not accent))
        y += 0.42
    return group(s, since(s, mark), head)


card(*cols[0], "Your request", "file-text", [
    ("circle-check", "424242", "Task: launch email"),
    ("circle-help", "E91D63", "Audience: ?"),
    ("circle-help", "E91D63", "Deadline: ?"),
])
card(*cols[1], "Its questions", "message-circle-question", [
    ("circle-help", "424242", "Who is it for?"),
    ("circle-help", "424242", "When does it end?"),
])
card(*cols[2], "Complete brief", "circle-check", [
    ("circle-check", "E91D63", "Task: launch email"),
    ("circle-check", "E91D63", "Audience: customers"),
    ("circle-check", "E91D63", "Deadline: Friday"),
], accent=True)
for ax in (3.33, 9.69 - 0.03 - 2.96 - 0.31):
    s.arrow(ax, TOP + CH / 2 - 0.06, 0.26, 0.12)
s.caption(
    0.34, 4.12, 9.32,
    "Ask only questions that could change the output. For minor gaps, state "
    "the assumption and carry on.",
)

# Why it matters: the same job on two timelines.
s = d.why("one short pause can prevent a full rewrite")
X0, UNIT, H = 2.70, 1.00, 0.46


def track(y, head, ico, fill, colour, segments):
    mark = len(s.raw.shapes)
    badge(s, ico, 0.34, y + (H - 0.42) / 2, 0.42, fill, colour,
          line=None if fill == ACCENT else RULE)
    s.text(0.86, y, 1.80, H, title_case(head), size=13, color=INK, bold=True,
           font=BODY_FONT, anchor="m")
    x = X0
    for label, w, style in segments:
        f, ln, tc = {"grey": (NODE_BG, RULE, INK), "tint": (NODE_ACC, ACCENT, ACCENT),
                     "pink": (ACCENT, ACCENT, "FFFFFF"),
                     "warn": ("FFFFFF", ACCENT, ACCENT)}[style]
        s.rect(x, y, w * UNIT - 0.06, H, fill=f, rounded=True, radius=0.08,
               line=ln, line_w=0.75)
        s.text(x, y, w * UNIT - 0.06, H, title_case(label), size=SZ_CAPTION,
               color=tc, bold=True, font=BODY_FONT, align="c", anchor="m")
        x += w * UNIT
    s.text(x + 0.02, y, 0.70, H, "Done", size=SZ_CAPTION, color=SOFT_INK,
           bold=True, font=BODY_FONT, anchor="m")
    return group(s, since(s, mark), head)


track(1.66, "Answer now", "rotate-ccw", NODE_BG, "424242", [
    ("Draft 1", 1.4, "grey"), ("Wrong audience", 1.9, "warn"),
    ("Rewrite", 1.3, "grey"), ("Draft 2", 1.3, "grey"),
])
track(2.40, "Ask first", "circle-help", ACCENT, "FFFFFF", [
    ("Questions", 1.2, "tint"), ("Answers", 1.1, "tint"), ("Draft 1", 1.5, "pink"),
])
# A time axis under both tracks.
mark = len(s.raw.shapes)
s.rect(X0, 3.12, 9.66 - X0, 0.012, fill=SOFT_INK)
s.arrow(9.40, 3.07, 0.26, 0.11)
icon(s, "clock", X0, 3.22, 0.22, "424242")
s.text(X0 + 0.30, 3.20, 2.00, 0.26, "Time", size=SZ_CAPTION, color=SOFT_INK,
       font=BODY_FONT, anchor="m")
group(s, since(s, mark), "Time axis")
# The honest failure mode.
mark = len(s.raw.shapes)
s.rect(0.34, 3.80, 9.32, 0.56, fill="FFFFFF", rounded=True, radius=0.08,
       line=RULE, line_w=0.75)
badge(s, "triangle-alert", 0.48, 3.87, 0.42, NODE_ACC, "E91D63")
s.text(1.04, 3.80, 8.50, 0.56,
       "Failure mode: too many low-value questions. Cap it at three.",
       size=SZ_CAPTION, color=INK, font=BODY_FONT, anchor="m")
group(s, since(s, mark), "Failure mode")

d.handoff(
    [
        ("Give a vague task", "leave one real gap"),
        ("Add the check", "ask before answering"),
        ("Answer it", "only what matters"),
        ("Compare", "fewer assumptions"),
    ],
    lead="Now we tell ChatGPT what to do when context is missing.",
)

path = d.save(out)
print(f"wrote {path}  ({len(d.prs.slides._sldIdLst)} slides)")
