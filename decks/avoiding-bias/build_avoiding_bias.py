#!/usr/bin/env python3
"""Build the seven-slide principles deck for recognising bias in a chat."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / ".claude" / "skills" / "principles-deck" / "scripts"))
from pptx.util import Inches, Pt  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.enum.dml import MSO_LINE  # noqa: E402
from deckkit import (Deck, title_case, INK, SOFT_INK, ACCENT, RULE, ARROW,  # noqa: E402
                     NODE_BG, NODE_ACC, BODY_FONT, SZ_CAPTION, SZ_DIAGRAM, _rgb)

ASSETS = Path(__file__).resolve().parent / "assets"


def mark(s):
    return len(s.raw.shapes)


def group(s, since, name):
    """Group every shape drawn since `since`, so it moves as one unit."""
    shapes = list(s.raw.shapes)[since:]
    if len(shapes) > 1:
        g = s.raw.shapes.add_group_shape(shapes)
        g.name = name
        return g


def oval(s, x, y, d, fill, line=None, dash=False):
    o = s.raw.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d),
                               Inches(d))
    o.fill.solid()
    o.fill.fore_color.rgb = _rgb(fill)
    if line:
        o.line.color.rgb = _rgb(line)
        o.line.width = Pt(0.75)
    else:
        o.line.fill.background()
    o.shadow.inherit = False
    return o


def badge(s, icon, x, y, d=0.42, fill="FFFFFF", line=RULE, tint="ink"):
    """A Lucide line icon (ISC licence) centred in a circle."""
    oval(s, x, y, d, fill, line)
    g = d * 0.56
    s.raw.shapes.add_picture(str(ASSETS / f"icon-{icon}-{tint}.png"),
                             Inches(x + (d - g) / 2), Inches(y + (d - g) / 2),
                             Inches(g), Inches(g))


def wire(s, x1, y1, x2, y2):
    """A straight connector line in the arrow pink."""
    c = s.raw.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = _rgb(ARROW)
    c.line.width = Pt(1.0)
    return c


def dashed_box(s, x, y, w, h):
    """The 'searched' area: a dashed accent outline, no fill."""
    r = s.rect(x, y, w, h, rounded=True, radius=0.08, line=ACCENT, line_w=1.5)
    r.line.dash_style = MSO_LINE.DASH
    return r

def label_table(data):
    """Header row and first column are labels, so they are Title Case."""
    return [[title_case(c) if (r == 0 or i == 0) else c for i, c in enumerate(row)]
            for r, row in enumerate(data)]


out = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parent / "avoiding-bias.pptx")
d = Deck()

d.title(
    "Avoiding Bias in a Chat",
    "How the context around a question can steer the answer before it is written.",
)

s = d.what("four sources can shape the answer")
sources = [
    ("brain", "Memory", "saved facts"),
    ("sliders-horizontal", "Custom instructions", "standing directions"),
    ("history", "Chat history", "earlier turns"),
    ("message-circle-question", "Your question", "the frame used now"),
]
cw, ch, bus_x = 3.70, 0.60, 4.44
centres = []
for i, (icon, label, note) in enumerate(sources):
    y = 1.56 + i * 0.72
    now = i == len(sources) - 1
    start = mark(s)
    s.rect(0.34, y, cw, ch, fill=NODE_ACC if now else NODE_BG, rounded=True,
           radius=0.05, line=ACCENT if now else RULE, line_w=1.0 if now else 0.75)
    badge(s, icon, 0.46, y + 0.09, fill="FFFFFF", line=ACCENT if now else RULE,
          tint="pink" if now else "ink")
    s.text(1.02, y + 0.07, cw - 0.80, 0.26, title_case(label), size=SZ_CAPTION,
           color=INK, bold=True, font=BODY_FONT)
    s.text(1.02, y + 0.31, cw - 0.80, 0.24, title_case(note), size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT)
    group(s, start, label)
    cy = y + ch / 2
    centres.append(cy)
    wire(s, 0.34 + cw, cy, bus_x, cy)
mid = sum(centres) / len(centres)
wire(s, bus_x, centres[0], bus_x, centres[-1])
s.arrow(bus_x, mid - 0.05, 0.40, 0.10)
start = mark(s)
s.rect(4.94, mid - 0.55, 1.90, 1.10, fill="FFFFFF", rounded=True, radius=0.06,
       line=RULE, line_w=0.75)
badge(s, "cpu", 4.94 + 0.74, mid - 0.44, fill=NODE_BG)
s.text(4.94, mid + 0.06, 1.90, 0.30, title_case("The model"), size=SZ_DIAGRAM,
       color=INK, bold=True, font=BODY_FONT, align="c")
group(s, start, "The model")
s.arrow(6.94, mid - 0.05, 0.40, 0.10)
s.node(7.44, mid - 0.55, 2.22, 1.10, "The answer", sub="tilted by all four",
       accent=True)
s.caption(0.34, 4.50, 9.32, "All four are useful. All four can tilt the answer.")

s = d.how("question framing changes what gets searched for")
# Each dot is a piece of evidence: pink supports the plan, grey challenges it.
# The dashed outline is where the question sends the search.
cards = [
    ("Question A", "\"Which evidence supports this plan?\"", False,
     "Finds one side only."),
    ("Question B", "\"What evidence supports or challenges this plan?\"", True,
     "Finds both sides."),
]
cw = 4.55
for i, (label, quote, both, result) in enumerate(cards):
    x = 0.34 + i * (cw + 0.22)
    start = mark(s)
    s.rect(x, 1.52, cw, 3.02, fill="FFFFFF", rounded=True, radius=0.05,
           line=ACCENT if both else RULE, line_w=1.0 if both else 0.75)
    s.text(x + 0.24, 1.64, cw - 0.48, 0.24, title_case(label), size=SZ_CAPTION,
           color=ACCENT if both else SOFT_INK, bold=True, font=BODY_FONT)
    s.text(x + 0.24, 1.90, cw - 0.48, 0.56, quote, size=SZ_DIAGRAM, color=INK,
           font=BODY_FONT, spacing=1.1)
    fx, fy, step, dot = x + 0.62, 2.72, 0.56, 0.30
    for col in range(6):
        for row in range(2):
            oval(s, fx + col * step, fy + row * 0.46, dot,
                 ACCENT if col < 3 else "C9C9CF")
    s.text(fx - 0.10, fy + 1.00, 3 * step, 0.24, title_case("for"),
           size=SZ_CAPTION, color=ACCENT, bold=True, font=BODY_FONT, align="c")
    s.text(fx + 3 * step - 0.10, fy + 1.00, 3 * step, 0.24, title_case("against"),
           size=SZ_CAPTION, color=SOFT_INK, bold=True, font=BODY_FONT, align="c")
    span = (6 if both else 3) * step - (step - dot)
    dashed_box(s, fx - 0.14, fy - 0.14, span + 0.28, 0.46 + dot + 0.28)
    s.text(x + 0.24, 4.14, cw - 0.48, 0.26, result, size=SZ_CAPTION,
           color=ACCENT if both else INK, bold=True, font=BODY_FONT)
    group(s, start, label)
s.caption(0.34, 4.66, 9.32, "The wording sets where the model looks.")

# Ownership words: models are tuned to please, so "I" and "my" invite agreement.
s = d.content("ownership words invite agreement")
mw, gap, top = 4.55, 0.22, 1.56
panels = []
for i, (tag, user, reply) in enumerate([
    ("Loaded",
     ["I wrote this launch plan and I", "think it's strong. Is it good?"],
     ["Yes, it's a strong plan. The", "timeline is realistic and the", "goals are clear."]),
    ("Neutral",
     ["Here is a launch plan. List its", "three biggest weaknesses."],
     ["1. No owner for customer support.", "2. The timeline has no buffer.",
      "3. Success is not measurable."]),
]):
    m = s.mock(0.34 + i * (mw + gap), top, mw, 3.00, app="AI Assistant", tag=tag)
    m.user(*user)
    with m.bubble():
        m.label("AI Assistant")
        for ln in reply:
            m.line(ln, highlight=True)
    m.fit()
    panels.append(m)
# Match the two windows so the comparison reads as a pair.
tallest = max(p.h for p in panels)
for p in panels:
    p.h = tallest
    p._panel.height = int(tallest * 914400)
# Name what each framing invites, under its window.
for i, (icon, tint, label) in enumerate([("triangle-alert", "ink", "Invites agreement"),
                                         ("check", "pink", "Invites critique")]):
    x = 0.34 + i * (mw + gap)
    start = mark(s)
    badge(s, icon, x, top + tallest + 0.12, d=0.34,
          fill=NODE_ACC if i else NODE_BG, line=ACCENT if i else RULE, tint=tint)
    s.text(x + 0.44, top + tallest + 0.12, 2.6, 0.34, title_case(label),
           size=SZ_CAPTION, color=ACCENT if i else INK, bold=True,
           font=BODY_FONT, anchor="m")
    group(s, start, label)
s.caption(
    0.34, top + tallest + 0.56, 9.32,
    "Models are tuned to please. \"I\" and \"my\" tell them what pleases you.",
)

# The fix: neutral wording, options side by side, ask for the case against.
s = d.content("ask neutral questions")
pairs = [
    ("\"I think X is better. Right?\"", "\"Which is better for <goal>: X or Y?\""),
    ("\"My essay is good, isn't it?\"", "\"Review this essay against <criteria>.\""),
    ("\"Why is X the best option?\"", "\"Compare X, Y and Z on <criteria>.\""),
    ("\"I'm sure the bug is in the API.\"", "\"Rank the likely causes of this bug.\""),
]
lx, lw, rx, rw = 0.34, 4.10, 5.10, 4.56
s.text(lx, 1.52, lw, 0.28, title_case("Loaded"), size=SZ_CAPTION, color=SOFT_INK,
       bold=True, font=BODY_FONT)
s.text(rx, 1.52, rw, 0.28, title_case("Neutral"), size=SZ_CAPTION, color=ACCENT,
       bold=True, font=BODY_FONT)
y, rh = 1.86, 0.56
for loaded, neutral in pairs:
    start = mark(s)
    s.rect(lx, y, lw, rh, fill=NODE_BG, rounded=True, radius=0.08, line=RULE,
           line_w=0.75)
    badge(s, "triangle-alert", lx + 0.10, y + 0.11, d=0.34, fill="FFFFFF")
    s.text(lx + 0.54, y, lw - 0.64, rh, loaded, size=SZ_CAPTION, color=SOFT_INK,
           font=BODY_FONT, anchor="m")
    s.arrow(lx + lw + 0.14, y + rh / 2 - 0.05, rx - lx - lw - 0.28, 0.10)
    s.rect(rx, y, rw, rh, fill=NODE_ACC, rounded=True, radius=0.08, line=ACCENT,
           line_w=1.0)
    badge(s, "check", rx + 0.10, y + 0.11, d=0.34, fill="FFFFFF", line=ACCENT,
          tint="pink")
    s.text(rx + 0.54, y, rw - 0.64, rh, neutral, size=SZ_CAPTION, color=INK,
           font=BODY_FONT, anchor="m")
    group(s, start, "Loaded to neutral")
    y += rh + 0.14
s.caption(0.34, y + 0.06, 9.32,
          "Drop \"I\" and \"my\". Put the options side by side. Ask for the case against.")

s = d.why("awareness gives you a way to check the frame")
checks = [
    ("brain", "Memory", "Preferences should persist",
     "Is an old preference steering this?"),
    ("sliders-horizontal", "Instructions", "The same rules should repeat",
     "Do the rules favour one outcome?"),
    ("history", "Chat history", "The task spans several turns",
     "Would a fresh chat answer differently?"),
    ("message-circle-question", "Question", "You need a focused answer",
     "Did the wording exclude alternatives?"),
]
gap = 0.20
cw = (9.32 - 3 * gap) / 4
for i, (icon, label, useful, check) in enumerate(checks):
    x = 0.34 + i * (cw + gap)
    start = mark(s)
    s.rect(x, 1.56, cw, 2.96, fill="FFFFFF", rounded=True, radius=0.05,
           line=RULE, line_w=0.75)
    badge(s, icon, x + 0.20, 1.72, d=0.46, fill=NODE_BG)
    s.text(x + 0.20, 2.30, cw - 0.40, 0.30, title_case(label), size=14, color=INK,
           bold=True, font=BODY_FONT)
    s.text(x + 0.20, 2.62, cw - 0.40, 0.60, useful, size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT, spacing=1.1)
    s.rect(x + 0.20, 3.28, cw - 0.40, 0.008, fill=RULE)
    s.text(x + 0.20, 3.38, cw - 0.40, 0.24, title_case("check"), size=SZ_CAPTION,
           color=ACCENT, bold=True, font=BODY_FONT)
    s.text(x + 0.20, 3.64, cw - 0.40, 1.00, check, size=SZ_CAPTION, color=INK,
           font=BODY_FONT, spacing=1.1)
    group(s, start, label)

d.handoff(
    [
        ("Inspect the context", "memory and instructions"),
        ("Open a fresh chat", "remove earlier framing"),
        ("Rewrite the question", "drop 'I' and 'my'"),
        ("Compare the answers", "notice what changed"),
    ],
    lead="See the frame, then test whether it changed the answer.",
)

path = d.save(out)
print(f"wrote {path}  ({len(d.prs.slides._sldIdLst)} slides)")
