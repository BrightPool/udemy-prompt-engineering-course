#!/usr/bin/env python3
"""Build the five-slide principles deck for specifying observable steps."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / ".claude" / "skills" / "principles-deck" / "scripts"))
from deckkit import Deck, title_case  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from deckkit import (ACCENT, ARROW, INK, SOFT_INK, RULE, NODE_BG, NODE_ACC,  # noqa: E402
                     BODY_FONT, HEAD_FONT, SZ_CAPTION, _rgb)

# Visual helpers for this deck (icons are Lucide, ISC licence, rendered to assets/).
ASSETS = Path(__file__).resolve().parent / "assets"
TONES = {"pink": (ACCENT, ACCENT, "white"), "soft": (NODE_ACC, ACCENT, "pink"),
         "grey": ("ECECF0", RULE, "ink")}


def mark(s):
    """Shape count now; pass it to group() to bundle everything drawn since."""
    return len(s.raw.shapes)


def group(s, first, name=None):
    """Group every shape drawn since `first`, so it moves as one unit."""
    shapes = s.raw.shapes
    g = shapes.add_group_shape(list(shapes)[first:])
    if name:
        g.name = name
    return g


def glyph(s, x, y, size, icon, colour="pink"):
    """A bare icon."""
    s.raw.shapes.add_picture(str(ASSETS / f"i-{icon}-{colour}.png"),
                             Inches(x), Inches(y), Inches(size), Inches(size))


def oval(s, x, y, d, fill, line=None):
    c = s.raw.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d),
                               Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = _rgb(fill)
    if line:
        c.line.color.rgb = _rgb(line)
        c.line.width = Pt(1.0)
    else:
        c.line.fill.background()
    c.shadow.inherit = False
    return c


def badge(s, x, y, d, icon, tone="pink"):
    """A round icon badge: pink (solid), soft (pink tint) or grey."""
    fill, line, colour = TONES[tone]
    oval(s, x, y, d, fill, line)
    g = d * 0.54
    glyph(s, x + (d - g) / 2, y + (d - g) / 2, g, icon, colour)

# The worked example: cars.csv (made up) enriched with real UK starting prices.
DATA = Path(__file__).resolve().parent / "data"


def cars_facts():
    """Headline numbers for the mock, read from the answer key."""
    import csv
    rows = list(csv.DictReader(open(DATA / "cars_enriched_answer_key.csv")))
    priced = [r for r in rows if r["price_gbp"]]
    by_body = {}
    for r in priced:
        by_body.setdefault(r["body_type"], []).append(float(r["price_gbp"]))
    avg = {k: sum(v) / len(v) for k, v in by_body.items()}
    top = sorted(priced, key=lambda r: -float(r["revenue_gbp"]))[:3]
    return {
        "n": len(rows), "priced": len(priced),
        "suv": avg.get("SUV", 0), "hatch": avg.get("Hatchback", 0),
        "top": [r["model"] for r in top],
    }


FACTS = cars_facts()


def k(v):
    """Pounds to the nearest thousand: 31,450 -> 31k."""
    return f"£{round(v / 1000)}k"


out = Path(sys.argv[1] if len(sys.argv) > 1 else
           Path(__file__).resolve().parent / "clear-instructions-steps.pptx")
d = Deck()

d.title(
    "Specifying the Steps",
    "How an ordered workflow helps ChatGPT complete tasks with real dependencies.",
)

# What it is: one ChatGPT request, three ordered steps, each with an output.
s = d.what("each step names the output the next one needs")
m = s.mock(0.34, 1.52, 6.05, 3.56)
m.user(
    "Using cars.csv, complete these steps in order:",
    "1. Add price_gbp from understandingdata.com/.../car-prices/v1",
    "2. Average price by body type; revenue = units x price.",
    "3. Save cars_enriched.csv: top 3 by revenue, any gaps.",
)
with m.bubble():
    m.label("AI Assistant")
    m.line(f"Step 1 done: {FACTS['priced']} of {FACTS['n']} rows priced.")
    m.line(f"Step 2 done: SUVs average {k(FACTS['suv'])}, "
           f"hatchbacks {k(FACTS['hatch'])}.")
    m.line("Step 3 done: saved. Top: " + ", ".join(FACTS["top"]) + ".",
           highlight=True)
    m.line("Gaps: none.", highlight=True)
m.fit()
# Beside the mock: the chain of outputs, one icon per step.
px = 6.78
chain = [("search", "Prices", "one per model"),
         ("calculator", "Analysis", "averages and revenue"),
         ("file-spreadsheet", "Enriched CSV", "plus a summary")]
for i, (icon, label, sub) in enumerate(chain):
    y = 1.60 + i * 0.98
    first = mark(s)
    badge(s, px, y, 0.52, icon, "pink" if i == 2 else "soft")
    s.text(px + 0.66, y - 0.02, 2.30, 0.28, title_case(label), size=14, color=INK,
           bold=True, font=BODY_FONT)
    s.text(px + 0.66, y + 0.26, 2.30, 0.28, title_case(sub), size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT)
    group(s, first, f"Output {i + 1}")
    if i < len(chain) - 1:
        s.arrow(px + 0.20, y + 0.58, 0.12, 0.32, direction="down")
s.caption(px, 4.52, 2.88, "Each output feeds the next.")

# How it works: three stations, the input file, and a check at each handover.
s = d.how("the output of one step becomes the next input")
col = 9.32 / 3
stations = [("search", "Price lookup", "one price per model", "Every row priced?",
             "file-spreadsheet", "Input: cars.csv"),
            ("calculator", "Analysis", "averages and revenue", "Sums add up?",
             "table-2", "Input: the priced rows"),
            ("file-spreadsheet", "Enriched file", "plus a short summary", "Gaps named?",
             "calculator", "Input: the analysis")]
bd, by = 0.80, 2.14
for i, (icon, label, sub, check, in_icon, in_text) in enumerate(stations):
    cx = 0.34 + i * col + col / 2
    first = mark(s)
    # What this stage consumes: the dependency, made visible.
    iw = 2.70
    s.rect(cx - iw / 2, 1.54, iw, 0.34, fill="FFFFFF", rounded=True, radius=0.17,
           line=RULE, line_w=0.75)
    glyph(s, cx - iw / 2 + 0.12, 1.59, 0.24, in_icon, "ink")
    s.text(cx - iw / 2 + 0.42, 1.54, iw - 0.50, 0.34, in_text, size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT, anchor="m")
    s.arrow(cx - 0.06, 1.90, 0.12, 0.20, direction="down")
    badge(s, cx - bd / 2, by, bd, icon, "pink" if i == 2 else "soft")
    s.text(cx - col / 2, by + 0.90, col, 0.30, title_case(label), size=14, color=INK,
           bold=True, font=BODY_FONT, align="c")
    s.text(cx - col / 2, by + 1.20, col, 0.28, title_case(sub), size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT, align="c")
    pw = 2.50
    s.rect(cx - pw / 2, by + 1.66, pw, 0.44, fill=NODE_ACC, rounded=True, radius=0.20,
           line=ACCENT, line_w=0.75)
    glyph(s, cx - pw / 2 + 0.16, by + 1.76, 0.24, "circle-check", "pink")
    s.text(cx - pw / 2 + 0.46, by + 1.66, pw - 0.56, 0.44, title_case(check),
           size=SZ_CAPTION, color=ACCENT, bold=True, font=BODY_FONT, anchor="m")
    group(s, first, f"Step {i + 1}")
    if i < len(stations) - 1:
        s.arrow(cx + bd / 2 + 0.16, by + bd / 2 - 0.06, col - bd - 0.32, 0.12)
s.caption(0.34, 4.46, 9.32,
          "If step 1 misses a price, the summary in step 3 should say so.")

# Why it matters: two lanes, a hidden path against a visible one.
s = d.why("a required process becomes easier to inspect")
lanes = [
    ("Outcome only", "Best for simple tasks", "A key stage is skipped", 1.62, False),
    ("Specified workflow", "Best for dependencies", "Rigid steps block a better route", 3.14, True),
]
for label, best, risk, y, steps in lanes:
    first = mark(s)
    s.text(0.34, y, 2.30, 0.30, title_case(label), size=14,
           color=ACCENT if steps else INK, bold=True, font=BODY_FONT)
    s.text(0.34, y + 0.30, 2.30, 0.26, title_case(best), size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT)
    badge(s, 2.70, y, 0.52, "message-square", "grey")
    badge(s, 9.14, y, 0.52, "circle-check", "pink" if steps else "grey")
    if steps:
        xs = [4.10, 5.75, 7.40]
        s.arrow(3.30, y + 0.20, xs[0] - 3.30 - 0.10, 0.12)
        for k, x in enumerate(xs):
            s.disc(x, y + 0.04, 0.44, str(k + 1), size=15)
            s.text(x - 0.40, y + 0.54, 1.24, 0.26, "Check", size=SZ_CAPTION,
                   color=ACCENT, bold=True, font=BODY_FONT, align="c")
            end = xs[k + 1] if k + 1 < len(xs) else 9.14
            s.arrow(x + 0.52, y + 0.20, end - x - 0.62, 0.12)
    else:
        s.arrow(3.30, y + 0.20, 0.44, 0.12)
        s.rect(3.86, y - 0.02, 4.64, 0.56, fill="EDEDF0", rounded=True, radius=0.10,
               line=RULE, line_w=0.75)
        s.text(3.86, y - 0.02, 4.64, 0.56, "?", size=22, color=SOFT_INK, bold=True,
               font=HEAD_FONT, align="c", anchor="m")
        s.arrow(8.58, y + 0.20, 0.46, 0.12)
    ry = y + (0.92 if steps else 0.72)
    glyph(s, 2.70, ry + 0.02, 0.22, "circle-alert", "ink")
    s.text(3.00, ry, 6.60, 0.26, "Risk: " + risk.lower() + ".", size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT)
    group(s, first, label)

d.handoff(
    [
        ("Upload cars.csv", "20 models, no prices"),
        ("Run the three steps", "prices, analysis, file"),
        ("Check the report", "every row priced"),
    ],
    lead="Now we run this three-step request live in ChatGPT on cars.csv.",
)

path = d.save(out)
print(f"wrote {path}  ({len(d.prs.slides._sldIdLst)} slides)")
