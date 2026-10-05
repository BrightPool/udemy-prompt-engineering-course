#!/usr/bin/env python3
"""Build the principles deck for What is Codex and /goal."""
import sys
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / ".claude" / "skills" / "principles-deck" / "scripts"))
from deckkit import (Deck, ACCENT, INK, SOFT_INK, RULE, NODE_ACC,  # noqa: E402
                     NODE_BG, BODY_FONT, SZ_CAPTION, title_case, _rgb)


def keep_case(slide, text):
    """deckkit Title Cases node labels; put back text that must keep its case."""
    cased = title_case(text)
    for shp in slide.raw.shapes:
        if shp.has_text_frame:
            for para in shp.text_frame.paragraphs:
                for run in para.runs:
                    if run.text == cased:
                        run.text = text


def badge(s, x, y, d, icon, fill=NODE_ACC, line=ACCENT, glyph=0.55):
    """A circle with an icon centred in it (Lucide icons, ISC licence)."""
    c = s.raw.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                               Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = _rgb(fill)
    c.line.color.rgb = _rgb(line)
    c.line.width = Pt(1.0)
    c.shadow.inherit = False
    g = d * glyph
    s.raw.shapes.add_picture(str(ASSETS / icon), Inches(x + (d - g) / 2),
                             Inches(y + (d - g) / 2), Inches(g), Inches(g))


def group(s, start, name):
    """Group every shape added since `start`, so it moves as one unit."""
    shapes = list(s.raw.shapes)[start:]
    if len(shapes) > 1:
        s.raw.shapes.add_group_shape(shapes).name = name


def logo(s, png, x, y, size):
    s.raw.shapes.add_picture(str(ASSETS / png), Inches(x), Inches(y),
                             Inches(size), Inches(size))

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
out = Path(sys.argv[1] if len(sys.argv) > 1 else HERE / "codex-goal.pptx")
d = Deck()

d.title(
    "What is Codex and /goal",
    "An agent that works in your files, and a way to hand it an objective instead of a prompt.",
)

# -- What is it: who runs each step of the loop -------------------------------
s = d.what("Codex does the work, not just the talking")
YOU, AI = "you", "ai"
loops = [
    ("icon-chatgpt.png", "ChatGPT", [
        (YOU, "icon-user-424242.png", "You ask"),
        (AI, "icon-bot-FFFFFF.png", "It answers"),
        (YOU, "icon-user-424242.png", "You copy and run it"),
        (YOU, "icon-user-424242.png", "You paste errors back"),
    ]),
    ("icon-codex-light.png", "Codex", [
        (YOU, "icon-user-424242.png", "You set the task"),
        (AI, "icon-file-search-FFFFFF.png", "It reads the files"),
        (AI, "icon-play-FFFFFF.png", "It edits and tests"),
        (AI, "icon-pen-line-FFFFFF.png", "It hands back changes"),
    ]),
]
px, arrow_w = 1.62, 0.26
pw = (9.66 - px - 3 * arrow_w) / 4
for r, (png, name, steps) in enumerate(loops):
    y = 1.66 + r * 1.12
    start = len(s.raw.shapes)
    logo(s, png, 0.34, y, 0.62)
    s.text(0.30, y + 0.66, 0.70, 0.26, name, size=SZ_CAPTION, color=INK, bold=True,
           font=BODY_FONT, align="c")
    for i, (who, icon, label) in enumerate(steps):
        x = px + i * (pw + arrow_w)
        ai = who == AI
        s.rect(x, y + 0.06, pw, 0.72, fill=ACCENT if ai else NODE_BG, rounded=True,
               radius=0.10, line=None if ai else RULE, line_w=0.75)
        s.raw.shapes.add_picture(str(ASSETS / icon), Inches(x + 0.14),
                                 Inches(y + 0.06 + 0.22), Inches(0.28), Inches(0.28))
        s.text(x + 0.52, y + 0.06, pw - 0.60, 0.72, title_case(label), size=SZ_CAPTION,
               color="FFFFFF" if ai else INK, bold=ai, font=BODY_FONT, anchor="m",
               spacing=1.1)
        if i < 3:
            s.arrow(x + pw + 0.04, y + 0.37, arrow_w - 0.08, 0.10)
    group(s, start, f"{name} loop")
# legend
ly = 3.98
start = len(s.raw.shapes)
s.rect(px, ly + 0.05, 0.22, 0.22, fill=NODE_BG, rounded=True, line=RULE, line_w=0.75)
s.text(px + 0.32, ly, 1.20, 0.32, "You", size=SZ_CAPTION, color=INK, font=BODY_FONT,
       anchor="m")
s.rect(px + 1.10, ly + 0.05, 0.22, 0.22, fill=ACCENT, rounded=True)
s.text(px + 1.42, ly, 1.60, 0.32, "The AI", size=SZ_CAPTION, color=INK, font=BODY_FONT,
       anchor="m")
group(s, start, "Legend")
s.text(px, ly + 0.44, 9.66 - px, 0.30,
       "Same account. The difference is who runs the loop.", size=14, color=INK,
       font=BODY_FONT)

# -- When to reach for it: task -> the right tool ------------------------------
s = d.content("When to reach for Codex")
tasks = [
    ("icon-message-square-424242.png", "Explain an idea or draft an email",
     "One answer, nothing to run", "icon-chatgpt.png", "ChatGPT"),
    ("icon-bug-424242.png", "Fix a bug or add a feature",
     "It edits code and runs tests", "icon-codex-light.png", "Codex"),
    ("icon-files-424242.png", "Build a report from 40 files in a folder",
     "It can open every file itself", "icon-codex-light.png", "Codex"),
    ("icon-replace-all-424242.png", "Change the same thing across many files",
     "It works through them one by one", "icon-codex-light.png", "Codex"),
]
for i, (icon, task, why, png, tool) in enumerate(tasks):
    y = 1.60 + i * 0.70
    start = len(s.raw.shapes)
    badge(s, 0.34, y + 0.04, 0.50, icon, fill=NODE_BG, line=RULE)
    s.text(1.04, y, 5.40, 0.30, title_case(task), size=14, color=INK, bold=True,
           font=BODY_FONT)
    s.text(1.04, y + 0.30, 5.40, 0.26, why, size=SZ_CAPTION, color=SOFT_INK,
           font=BODY_FONT)
    s.arrow(6.62, y + 0.24, 0.62, 0.10)
    logo(s, png, 7.50, y + 0.03, 0.50)
    s.text(8.10, y + 0.03, 1.50, 0.50, tool, size=14, bold=True,
           color=ACCENT if tool == "Codex" else INK, font=BODY_FONT, anchor="m")
    if i < len(tasks) - 1:
        s.rect(0.34, y + 0.63, 9.32, 0.008, fill=RULE)
    group(s, start, f"Task {i + 1}")
s.caption(0.34, 4.50, 9.32,
          "Rule of thumb: if the work lives in files, give it to Codex.")

# -- How /goal works: the loop as a ring, the commands beside it ---------------
s = d.how("/goal keeps Codex working until it is done")
D, cx, cy = 2.50, 3.55, 3.28
ring = s.raw.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - D / 2), Inches(cy - D / 2),
                              Inches(D), Inches(D))
ring.fill.background()
ring.line.color.rgb = _rgb(ACCENT)
ring.line.width = Pt(2.0)
ring.shadow.inherit = False
logo(s, "icon-codex-light.png", cx - 0.40, cy - 0.52, 0.80)
s.text(cx - 0.90, cy + 0.30, 1.80, 0.26, title_case("for hours if needed"),
       size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT, align="c")
bd = 0.56
stations = [  # (angle position, icon, label, sub, label side)
    ((cx, cy - D / 2), "icon-map-E91D63.png", "Plan", "the next step", "n"),
    ((cx + D / 2, cy), "icon-hammer-E91D63.png", "Act", "edit, run, build", "r"),
    ((cx, cy + D / 2), "icon-list-checks-E91D63.png", "Check", "against the goal", "n"),
    ((cx - D / 2, cy), "icon-flag-E91D63.png", "Done?", "stop or go again", "l"),
]
for (bx, by), icon, lab, sub, side in stations:
    start = len(s.raw.shapes)
    badge(s, bx - bd / 2, by - bd / 2, bd, icon, fill="FFFFFF")
    # labels sit clear of the ring: top/bottom stations step out past the arc
    if side == "r":
        tx, al = bx + bd / 2 + 0.10, "l"
    elif side == "n":
        tx, al = bx + 0.95, "l"
    else:
        tx, al = bx - bd / 2 - 1.80, "r"
    s.text(tx, by - 0.28, 1.70, 0.28, lab, size=14, color=INK, bold=True,
           font=BODY_FONT, align=al)
    s.text(tx, by, 1.70, 0.26, title_case(sub), size=SZ_CAPTION, color=SOFT_INK,
           font=BODY_FONT, align=al, wrap=False)
    group(s, start, lab)
# commands, typed exactly
kx = 7.05
s.text(kx, 1.56, 2.61, 0.30, title_case("The commands"), size=14, color=INK,
       bold=True, font=BODY_FONT)
cmds = [("/goal <objective>", "start"), ("/goal", "see progress"),
        ("/goal pause", "stop for now"), ("/goal resume", "carry on"),
        ("/goal clear", "end it")]
for i, (cmd, what) in enumerate(cmds):
    y = 1.96 + i * 0.58
    start = len(s.raw.shapes)
    s.text(kx, y, 2.61, 0.26, cmd, size=13, color=ACCENT, bold=True, font=BODY_FONT)
    s.text(kx, y + 0.24, 2.61, 0.24, title_case(what), size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT)
    s.rect(kx, y + 0.52, 2.61, 0.008, fill=RULE)
    group(s, start, cmd)

# -- Prompt vs goal: two timelines, then the four parts of a goal ------------
s = d.content("A prompt steers one turn, a goal defines done")
tx0 = 1.70
s.text(0.34, 1.66, 1.30, 0.36, title_case("A prompt"), size=14, color=INK, bold=True,
       font=BODY_FONT, anchor="m")
start = len(s.raw.shapes)
for i in range(8):
    you = i % 2 == 0
    x = tx0 + i * 0.66
    s.rect(x, 1.68, 0.58, 0.32, fill=NODE_BG if you else NODE_ACC, rounded=True,
           radius=0.10, line=RULE if you else ACCENT, line_w=0.75)
    s.text(x, 1.68, 0.58, 0.32, "You" if you else "AI", size=SZ_CAPTION,
           color=INK if you else ACCENT, font=BODY_FONT, align="c", anchor="m")
s.text(tx0 + 8 * 0.66 + 0.10, 1.66, 2.40, 0.36, title_case("you steer every turn"),
       size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT, anchor="m")
group(s, start, "Prompt timeline")
s.text(0.34, 2.36, 1.30, 0.36, title_case("A goal"), size=14, color=ACCENT, bold=True,
       font=BODY_FONT, anchor="m")
start = len(s.raw.shapes)
s.rect(tx0, 2.38, 0.58, 0.32, fill=NODE_BG, rounded=True, radius=0.10, line=RULE,
       line_w=0.75)
s.text(tx0, 2.38, 0.58, 0.32, "You", size=SZ_CAPTION, color=INK, font=BODY_FONT,
       align="c", anchor="m")
bar_x, bar_w = tx0 + 0.66, 7 * 0.66 - 0.08 - 0.44
s.rect(bar_x, 2.38, bar_w, 0.32, fill=ACCENT, rounded=True, radius=0.10)
s.text(bar_x, 2.38, bar_w, 0.32, "Codex: plan, act, check, repeat", size=SZ_CAPTION,
       color="FFFFFF", bold=True, font=BODY_FONT, align="c", anchor="m")
badge(s, bar_x + bar_w + 0.08, 2.36, 0.36, "icon-flag-FFFFFF.png", fill=ACCENT)
s.text(tx0 + 8 * 0.66 + 0.10, 2.36, 2.40, 0.36, title_case("runs to the stop line"),
       size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT, anchor="m")
group(s, start, "Goal timeline")
s.rect(0.34, 3.00, 9.32, 0.008, fill=RULE)
s.text(0.34, 3.12, 9.32, 0.28, title_case("Every goal has four parts"), size=13,
       color=INK, bold=True, font=BODY_FONT)
parts = [("icon-target-E91D63.png", "Objective", "what to achieve"),
         ("icon-lock-E91D63.png", "Limits", "no-go areas"),
         ("icon-list-checks-E91D63.png", "Proof", "how to check it"),
         ("icon-flag-E91D63.png", "Stop", "when to finish")]
pw = 9.32 / 4
for i, (icon, lab, sub) in enumerate(parts):
    x = 0.34 + i * pw
    start = len(s.raw.shapes)
    badge(s, x, 3.52, 0.50, icon)
    s.text(x + 0.62, 3.50, pw - 0.66, 0.28, lab, size=14, color=INK, bold=True,
           font=BODY_FONT)
    s.text(x + 0.62, 3.78, pw - 0.66, 0.26, title_case(sub), size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT)
    group(s, start, lab)
s.caption(0.34, 4.40, 9.32,
          "Can't say how to check it? Use a prompt and steer instead.")

# -- Anatomy of a good goal ---------------------------------------------------
s = d.content("Anatomy of a good goal")
m = s.mock(0.34, 1.52, 5.70, 3.50, app="Codex", tag="Example")
m.label("User")
m.para("/goal Turn the 40 interview transcripts in /interviews into a findings report.")
m.gap(0.06)
m.line("Read first: questions.md and every transcript.")
m.line("Don't change: the transcripts themselves.")
m.line("Prove it: each finding cites two or more transcripts.",
       highlight=True)
m.line("Stop when: all five questions are answered.",
       highlight=True)
m.line("And every quote matches its source file.", highlight=True)
m.composer(placeholder="Message Codex…", pin=False)
callouts = [("icon-target-E91D63.png", "Objective", "the /goal line"),
            ("icon-book-open-E91D63.png", "Read First", "where to start"),
            ("icon-lock-E91D63.png", "Limits", "what not to change"),
            ("icon-list-checks-E91D63.png", "Proof", "how to check each finding"),
            ("icon-flag-E91D63.png", "Stop", "when to finish")]
for i, (icon, lab, sub) in enumerate(callouts):
    y = 1.60 + i * 0.58
    start = len(s.raw.shapes)
    badge(s, 6.34, y, 0.44, icon)
    s.text(6.92, y - 0.04, 2.74, 0.26, lab, size=13, color=INK, bold=True,
           font=BODY_FONT)
    s.text(6.92, y + 0.20, 2.74, 0.26, title_case(sub), size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT)
    group(s, start, lab)
s.text(6.34, 4.56, 3.32, 0.46, "These let it run without you.",
       size=SZ_CAPTION, color=ACCENT, bold=True, font=BODY_FONT, spacing=1.1)

d.handoff(
    [
        ("Open a folder", "in Codex"),
        ("Set one goal", "with a stop line"),
        ("Watch the loop", "then check the work"),
    ],
    lead="Next lecture: a second AI reviews the work before you do.",
)

path = d.save(out)
print(f"wrote {path}  ({len(d.prs.slides._sldIdLst)} slides)")
