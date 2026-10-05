#!/usr/bin/env python3
"""Build the intro deck (Introduction to Core AI Skills) for the section Core Skills for Working with AI."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / ".claude" / "skills" / "principles-deck" / "scripts"))
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from deckkit import (Deck, ACCENT, INK, SOFT_INK, RULE, NODE_ACC,  # noqa: E402
                     BODY_FONT, SZ_CAPTION, SZ_DIAGRAM, _rgb)

TOP = 1.56

out = Path(sys.argv[1] if len(sys.argv) > 1 else
           Path(__file__).resolve().parent / "core-skills-intro.pptx")
d = Deck()

d.title(
    "Introduction to Core AI Skills",
    "The habits and techniques that consistently improve what AI gives you back.",
)

# The same request twice: vague, then with the core skills applied.
s = d.what("the same model, a better request")
mw, gap, top = 4.55, 0.22, 1.52
panels = []
for i, (tag, user, reply) in enumerate([
    ("Before",
     ["Write an email about the price change."],
     ["Dear customer, we are writing to", "inform you of some changes to our",
      "pricing. Please contact us with", "any questions."]),
    ("After",
     ["Act as our account manager. Email",
      "existing customers about the 10% rise",
      "on 1 March. Say why, keep it under",
      "120 words, end with a booking link."],
     ["Subject: Your plan from 1 March", "Your plan rises 10% to fund more support.",
      "Book a 15-minute review: [link]"]),
]):
    m = s.mock(0.34 + i * (mw + gap), top, mw, 3.30, app="AI Assistant", tag=tag)
    m.user(*user)
    with m.bubble():
        m.label("AI Assistant")
        for ln in reply:
            m.line(ln, highlight=(i == 1))
    m.fit(pad=0.14)
    panels.append(m)
# Match the two windows so the comparison reads as a pair.
tallest = max(p.h for p in panels)
for p in panels:
    p.h = tallest
    p._panel.height = int(tallest * 914400)
# Name what changed, under the "after" window.
chips = ["Role", "Context", "Format", "Length"]
cx, cy = 0.34 + mw + gap, top + tallest + 0.18
for chip in chips:
    cw = 0.30 + len(chip) * 0.105
    s.rect(cx, cy, cw, 0.30, fill=NODE_ACC, rounded=True, radius=0.15,
           line=ACCENT, line_w=0.75)
    s.text(cx, cy, cw, 0.30, chip, size=SZ_CAPTION, color=ACCENT, bold=True,
           font=BODY_FONT, align="c", anchor="m")
    cx += cw + 0.10
s.text(0.34, cy, mw, 0.30, "Same model. Only the request changed.",
       size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT, anchor="m")

# Improving AI Outputs, folded in: four levers drawn as nested rings, because
# each one wraps the one inside it (the model sees context, uses tools, inside a
# workflow).
s = d.content("Four levers you control")
rings = [  # outermost first: (label, fill, line, text colour)
    ("Workflow", "F5F5F7", RULE, INK),
    ("Tools", "ECECF0", RULE, INK),
    ("Context", NODE_ACC, ACCENT, ACCENT),
    ("Model", ACCENT, ACCENT, "FFFFFF"),
]
D, step = 3.36, 0.80
cx, cy = 0.34 + D / 2, TOP + D / 2
for i, (lab, fill, line, ink) in enumerate(rings):
    dia = D - i * step
    ring = s.raw.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - dia / 2),
                                  Inches(cy - dia / 2), Inches(dia), Inches(dia))
    ring.fill.solid()
    ring.fill.fore_color.rgb = _rgb(fill)
    ring.line.color.rgb = _rgb(line)
    ring.line.width = Pt(1.0)
    ring.shadow.inherit = False
    core = i == len(rings) - 1
    ly = cy - 0.14 if core else cy - dia / 2 + 0.07
    s.text(cx - 0.9, ly, 1.8, 0.28, lab, size=SZ_DIAGRAM, color=ink, bold=True,
           font=BODY_FONT, align="c", anchor="m")

# The key: inside out, what each lever is and where the section covers it.
kx, kw = 4.30, 5.36
key = [  # each lever gets its own icon (Lucide, ISC licence), in its ring colour
    ("Model", "Which one does the work", "Pick Per Task", ACCENT, ACCENT, "model"),
    ("Context", "What it sees", "Most of This Section", NODE_ACC, ACCENT, "context"),
    ("Tools", "What it can do", "Codex and /goal", "ECECF0", RULE, "tools"),
    ("Workflow", "The steps around it", "Adversarial Review", "F5F5F7", RULE,
     "workflow"),
]
assets = Path(__file__).resolve().parent / "assets"
ky, row = TOP + 0.08, 0.66
for lab, what, where, fill, line, icon in key:
    badge = 0.46
    dot = s.raw.shapes.add_shape(MSO_SHAPE.OVAL, Inches(kx), Inches(ky - 0.03),
                                 Inches(badge), Inches(badge))
    dot.fill.solid()
    dot.fill.fore_color.rgb = _rgb(fill)
    dot.line.color.rgb = _rgb(line)
    dot.line.width = Pt(1.0)
    dot.shadow.inherit = False
    glyph = 0.26
    s.raw.shapes.add_picture(str(assets / f"icon-{icon}.png"),
                             Inches(kx + (badge - glyph) / 2),
                             Inches(ky - 0.03 + (badge - glyph) / 2),
                             Inches(glyph), Inches(glyph))
    s.text(kx + 0.60, ky, 1.20, 0.40, lab, size=14, color=INK, bold=True,
           font=BODY_FONT, anchor="m")
    s.text(kx + 1.72, ky, kw - 1.72, 0.22, what, size=SZ_CAPTION, color=INK,
           font=BODY_FONT)
    s.text(kx + 1.72, ky + 0.22, kw - 1.72, 0.22, where, size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT)
    s.rect(kx, ky + row - 0.10, kw, 0.008, fill=RULE)
    ky += row
s.caption(
    kx, ky + 0.06, kw,
    "When an output disappoints, ask which lever to pull. Rewording the prompt "
    "is only one of them.",
)

# The route through the section, in recording order.
s = d.content("The route through this section")
route = [
    "Avoiding Bias", "Specifying the Steps", "Meta Prompting",
    "Transforming Content for Your Audience", "Role Prompting", "Ask for Context",
    "Key Phrases To Add To Your Prompts", "Different Output Formats",
    "Overcoming the Token Limit", "Least to Most", "Delimiters",
    "Pre-Warming Chats", "What is Codex and /goal", "Adversarial Review of Outputs",
]
col_w, row_h = 4.55, 0.44
shapes = s.raw.shapes
columns = [[], []]
for i, name in enumerate(route):
    first = len(shapes)
    x = 0.34 + (i // 7) * (col_w + 0.22)
    y = TOP + (i % 7) * row_h
    s.disc(x + 0.08, y + 0.07, 0.30, str(i + 1), size=12)
    s.text(x + 0.52, y, col_w - 0.60, row_h, name, size=SZ_CAPTION,
           color=INK, font=BODY_FONT, anchor="m",
           wrap=False)
    if i % 7 != 6:
        s.rect(x + 0.52, y + row_h - 0.005, col_w - 0.52, 0.008, fill=RULE)
    row = shapes.add_group_shape(list(shapes)[first:])
    row.name = f"Lecture {i + 1}"
    columns[i // 7].append(row)
for n, col in enumerate(columns, 1):
    shapes.add_group_shape(col).name = f"Lectures column {n}"

path = d.save(out)
print(f"wrote {path}  ({len(d.prs.slides._sldIdLst)} slides)")
