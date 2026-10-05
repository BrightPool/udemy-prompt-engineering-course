#!/usr/bin/env python3
"""Build the principles deck for Key Phrases To Add To Your Prompts."""
import sys
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / ".claude" / "skills" / "principles-deck" / "scripts"))
from deckkit import (Deck, title_case, _rgb, ACCENT, INK, SOFT_INK, RULE,  # noqa: E402
                     NODE_ACC, BODY_FONT, SZ_CAPTION)

ASSETS = Path(__file__).resolve().parent / "assets"   # Lucide icons, ISC licence


def badge(s, x, y, d, icon, fill=NODE_ACC, line=ACCENT, glyph=0.55):
    """A circle with an icon centred in it."""
    c = s.raw.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                               Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = _rgb(fill)
    c.line.color.rgb = _rgb(line)
    c.line.width = Pt(1.0)
    c.shadow.inherit = False
    g = d * glyph
    s.raw.shapes.add_picture(str(ASSETS / icon), Inches(x + (d - g) / 2),
                             Inches(y + (d - g) / 2), Inches(g), Inches(g))


def group(s, start, name):
    """Group every shape added since `start`, so it moves as one unit."""
    shapes = list(s.raw.shapes)[start:]
    if len(shapes) > 1:
        s.raw.shapes.add_group_shape(shapes).name = name


out = Path(sys.argv[1] if len(sys.argv) > 1 else
           Path(__file__).resolve().parent / "key-phrases.pptx")
d = Deck()

d.title(
    "Key Phrases To Add To Your Prompts",
    "Short, reusable phrases that set how much, how long and when to stop.",
)

# What it is: the phrases in a real prompt, each one named beside it.
s = d.what("a phrase that sets the finish line")
m = s.mock(0.40, 1.56, 4.60, 3.40, generic=True)
m.user("Turn these meeting notes into a follow-up email.")
m.line("Make it at least 3 sentences.", highlight=True)
m.line("List every action with an owner.")
m.line("Don't stop until each action is covered.", highlight=True)
m.composer(pin=False)
ax = 5.40
for i, (icon, name, what, example) in enumerate([
    ("icon-ruler-FFFFFF.png", "Floor", "A minimum it must reach",
     "\"at least 3 sentences\""),
    ("icon-flag-FFFFFF.png", "Objective", "A finish line it must cross",
     "\"until each action is covered\""),
]):
    y = 1.70 + i * 1.10
    start = len(s.raw.shapes)
    badge(s, ax, y, 0.62, icon, fill=ACCENT)
    s.text(ax + 0.82, y - 0.02, 3.40, 0.30, name, size=15, color=INK, bold=True,
           font=BODY_FONT)
    s.text(ax + 0.82, y + 0.30, 3.40, 0.26, title_case(what), size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT)
    s.text(ax + 0.82, y + 0.56, 3.40, 0.26, example, size=SZ_CAPTION, color=ACCENT,
           font=BODY_FONT)
    group(s, start, name)
s.text(ax, 4.00, 4.20, 0.50, "Without one, the model guesses when it is done.",
       size=14, color=INK, font=BODY_FONT, spacing=1.15)

# The four families, each with its icon and one example to copy.
s = d.content("Four phrase families")
families = [
    ("icon-ruler-E91D63.png", "Floors", "How Much",
     "\"at least 200 rows\""),
    ("icon-flag-E91D63.png", "Objectives", "When to Stop",
     "\"Don't stop until it is complete\""),
    ("icon-timer-E91D63.png", "Time", "How Long",
     "\"Work for as long as 2 hours\""),
    ("icon-scan-search-E91D63.png", "Review", "Who Checks",
     "\"Review it as a strict editor\""),
]
gap = 0.20
cw = (9.32 - 3 * gap) / 4
for i, (icon, name, controls, example) in enumerate(families):
    x = 0.34 + i * (cw + gap)
    start = len(s.raw.shapes)
    badge(s, x + (cw - 0.80) / 2, 1.60, 0.80, icon)
    s.text(x, 2.54, cw, 0.32, name, size=15, color=INK, bold=True,
           font=BODY_FONT, align="c")
    s.text(x, 2.88, cw, 0.26, controls, size=SZ_CAPTION, color=ACCENT, bold=True,
           font=BODY_FONT, align="c")
    # the example as a quote: a thin accent bar, not a box
    s.rect(x + 0.10, 3.36, 0.03, 0.50, fill=ACCENT)
    s.text(x + 0.22, 3.36, cw - 0.30, 0.56, example, size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT, spacing=1.15)
    group(s, start, name)
s.caption(0.34, 4.20, 9.32,
          "Each one answers a question the model would otherwise guess.")

# Why: vague phrases on the left, checkable rewrites on the right.
s = d.why("checkable beats vague")
lx, rx, top = 0.34, 5.20, 1.62
s.text(lx, top, 4.20, 0.30, title_case("Hard to check"), size=14, color=SOFT_INK,
       bold=True, font=BODY_FONT)
s.text(rx, top, 4.40, 0.30, title_case("Easy to check"), size=14, color=ACCENT,
       bold=True, font=BODY_FONT)
s.rect(lx, top + 0.38, 9.32, 0.012, fill=INK)
for i, (vague, clear) in enumerate([
    ("\"Make it good\"", "\"At least 3 sentences\""),
    ("\"Be thorough\"", "\"Cover every action item\""),
    ("\"Write a detailed sheet\"", "\"At least 200 rows\""),
]):
    y = top + 0.56 + i * 0.62
    start = len(s.raw.shapes)
    s.raw.shapes.add_picture(str(ASSETS / "icon-circle-x-6B6B6B.png"),
                             Inches(lx), Inches(y + 0.04), Inches(0.30),
                             Inches(0.30))
    s.text(lx + 0.46, y, 3.60, 0.38, vague, size=14, color=SOFT_INK,
           font=BODY_FONT, anchor="m")
    s.arrow(4.40, y + 0.14, 0.50, 0.10)
    s.raw.shapes.add_picture(str(ASSETS / "icon-circle-check-E91D63.png"),
                             Inches(rx), Inches(y + 0.04), Inches(0.30),
                             Inches(0.30))
    s.text(rx + 0.46, y, 4.00, 0.38, clear, size=14, color=INK, bold=True,
           font=BODY_FONT, anchor="m")
    s.rect(lx, y + 0.52, 9.32, 0.008, fill=RULE)
    group(s, start, f"Rewrite {i + 1}")
s.caption(
    0.34, 4.14, 9.32,
    "The model honours what it can measure. A floor can invite padding, and "
    "time phrases only matter in agents like Codex that keep working.",
)

# Cheat sheet: one row per family, icon beside the phrase to copy.
s = d.content("Cheat sheet: copy these")
rows = [
    ("icon-ruler-FFFFFF.png",
     "Make it at least <N> <sentences / rows / examples>."),
    ("icon-flag-FFFFFF.png",
     "Don't stop until you have completed: <objective>."),
    ("icon-timer-FFFFFF.png",
     "Work for as long as <N> hours, checking progress as you go."),
    ("icon-scan-search-FFFFFF.png",
     "Then review your work against <criteria> and fix what fails."),
]
for i, (icon, phrase) in enumerate(rows):
    y = 1.62 + i * 0.74
    start = len(s.raw.shapes)
    badge(s, 0.34, y, 0.50, icon, fill=ACCENT)
    s.text(1.10, y, 8.56, 0.50, phrase, size=14, color=INK, font=BODY_FONT,
           anchor="m")
    if i < len(rows) - 1:
        s.rect(1.10, y + 0.62, 8.56, 0.008, fill=RULE)
    group(s, start, f"Phrase {i + 1}")
s.caption(0.34, 4.66, 9.32,
          "Swap the <placeholders>. Keep the rest word for word.")

d.handoff(
    [
        ("Add a floor", "to a short email"),
        ("Set an objective", "and watch it finish"),
        ("Ask for a review", "a preview of a later lecture"),
    ],
    lead="Next, we add these phrases to real prompts and compare the results.",
)

path = d.save(out)
print(f"wrote {path}  ({len(d.prs.slides._sldIdLst)} slides)")
