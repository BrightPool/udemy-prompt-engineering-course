#!/usr/bin/env python3
"""
"Local Models with LM Studio" - the principles half of the re-film.

Run:
    uv run --with python-pptx python decks/local-models-lm-studio/build_local_models_lm_studio.py
    uv run --with python-pptx python .claude/skills/principles-deck/scripts/check_deck.py \
        decks/local-models-lm-studio/local-models-lm-studio.pptx

Arc: why you might run a model on your own machine, and what you give up (what) -
the local server speaks the same API, so only the address changes (how).

No model names, sizes or version numbers. The live demo picks today's model.
"""
import sys
from pathlib import Path
from types import SimpleNamespace

SKILL = Path(__file__).resolve().parents[2] / ".claude" / "skills" / "principles-deck"
sys.path.insert(0, str(SKILL / "scripts"))
from deckkit import (ACCENT, BODY_FONT, INK, RULE, SOFT_INK, SZ_CAPTION,  # noqa: E402
                     SZ_DIAGRAM, Deck, title_case)
from code_image import code_image  # noqa: E402
from pptx.util import Inches  # noqa: E402

ASSETS = Path(__file__).resolve().parent / "assets"
out = Path(sys.argv[1] if len(sys.argv) > 1 else
           Path(__file__).resolve().parent / "local-models-lm-studio.pptx")


def picture(s, path, x, y, w, ratio):
    h = w / ratio
    s.raw.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    return y + h


def panels(s, y, left_title, left, right_title, right, h, x=0.34, w=9.32,
           gap=0.40):
    """Two labelled panels, like deckkit's compare(), with a dot instead of a
    dash so no em dash reaches the slide."""
    col = (w - gap) / 2
    for cx, title, items, accent in ((x, left_title, left, False),
                                     (x + col + gap, right_title, right, True)):
        s.rect(cx, y, col, h, fill="FFFFFF", rounded=True,
               line=ACCENT if accent else RULE, line_w=1.0 if accent else 0.75)
        s.text(cx + 0.20, y + 0.16, col - 0.40, 0.30, title_case(title), size=14,
               color=ACCENT if accent else INK, bold=True, font=BODY_FONT)
        for i, item in enumerate(items):
            s.text(cx + 0.20, y + 0.60 + i * 0.34, col - 0.40, 0.30,
                   f"\u2022  {item}", size=SZ_DIAGRAM, color=INK, font=BODY_FONT)


d = Deck()
d.title("Local Models with LM Studio",
        "How to run an AI model on your own machine and call it like a cloud "
        "API.")

s = d.what("a model that runs on your own machine")
panels(s, 1.56,
          "Why Run It Locally", ["Prompts stay on your machine",
                                 "No per-token bill",
                                 "Works offline",
                                 "Free to experiment and repeat"],
          "What You Give Up", ["Smaller models, less capable",
                               "Limited by your RAM and GPU",
                               "You download and update models"],
          h=1.98)
s.caption(0.34, 3.74, 9.32,
          "Reach for local when privacy or volume matters more than peak quality.")

s = d.how("same SDK, different address")
png, ratio = code_image('''from openai import OpenAI

client = OpenAI(
    base_url="http://localhost:1234/v1",
    api_key="lm-studio",  # any placeholder
)''', out_dir=ASSETS, name="lm-studio-client")
bottom = picture(s, png, 0.40, 1.56, 5.30, ratio)
s.text(0.40, bottom + 0.12, 5.30, 0.30,
       "Change the address and key; name the local model.", size=SZ_CAPTION, color=SOFT_INK,
       font=BODY_FONT)
s.beside(SimpleNamespace(x=0.40, y=1.56, w=5.30, h=bottom - 1.56),
         "LM Studio serves the loaded model behind an OpenAI-compatible endpoint "
         "on your machine. Point the client at it and the rest of your code stays "
         "the same.")

d.handoff([("Load a model", "in LM Studio"),
           ("Start the server", "on localhost"),
           ("Change the base URL", "and run the notebook")],
          lead="That is the principle. Now we run it on this machine.")

path = d.save(out)
print(f"wrote {path}  ({len(d.prs.slides._sldIdLst)} slides)")
