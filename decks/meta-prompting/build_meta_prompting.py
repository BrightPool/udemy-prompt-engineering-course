#!/usr/bin/env python3
"""Build the five-slide principles deck for meta prompting.

Meta prompting here means: once a ChatGPT conversation has produced the output you
want, ask ChatGPT to turn that conversation into a reusable prompt, so the task can
be repeated with the same result. (Asking the model to question you first is the
separate "Ask for Context" lecture.)
"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / ".claude" / "skills" / "principles-deck" / "scripts"))
from deckkit import Deck, title_case  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from deckkit import (ACCENT, ARROW, INK, SOFT_INK, RULE, NODE_BG, NODE_ACC,  # noqa: E402
                     BODY_FONT, HEAD_FONT, SZ_CAPTION, _rgb)

# Visual helpers for this deck (icons are Lucide, ISC licence, rendered to assets/).
ASSETS = Path(__file__).resolve().parent / "assets"
TONES = {"pink": (ACCENT, ACCENT, "white"), "soft": (NODE_ACC, ACCENT, "pink"),
         "grey": ("ECECF0", RULE, "ink")}


def mark(s):
    """Shape count now; pass it to group() to bundle everything drawn since."""
    return len(s.raw.shapes)


def group(s, first, name=None):
    """Group every shape drawn since `first`, so it moves as one unit."""
    shapes = s.raw.shapes
    g = shapes.add_group_shape(list(shapes)[first:])
    if name:
        g.name = name
    return g


def glyph(s, x, y, size, icon, colour="pink"):
    """A bare icon."""
    s.raw.shapes.add_picture(str(ASSETS / f"i-{icon}-{colour}.png"),
                             Inches(x), Inches(y), Inches(size), Inches(size))


def oval(s, x, y, d, fill, line=None):
    c = s.raw.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d),
                               Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = _rgb(fill)
    if line:
        c.line.color.rgb = _rgb(line)
        c.line.width = Pt(1.0)
    else:
        c.line.fill.background()
    c.shadow.inherit = False
    return c


def badge(s, x, y, d, icon, tone="pink"):
    """A round icon badge: pink (solid), soft (pink tint) or grey."""
    fill, line, colour = TONES[tone]
    oval(s, x, y, d, fill, line)
    g = d * 0.54
    glyph(s, x + (d - g) / 2, y + (d - g) / 2, g, icon, colour)

out = Path(sys.argv[1] if len(sys.argv) > 1 else
           Path(__file__).resolve().parent / "meta-prompting.pptx")
d = Deck()

d.title(
    "Meta Prompting",
    "Turning a ChatGPT conversation that worked into a prompt you can reuse.",
)

# What it is: many corrections on the left collapse into one prompt on the right.
s = d.what("the conversation becomes the prompt")
s.text(0.34, 1.56, 3.05, 0.30, title_case("Ten turns of corrections"), size=13,
       color=INK, bold=True, font=BODY_FONT)
corrections = ["Cut it to 150 words.", "Lead with the benefit.",
               "Only three updates.", "Drop the jargon.", "End with one clear CTA."]
for i, text in enumerate(corrections):
    y = 1.96 + i * 0.50
    first = mark(s)
    s.rect(0.34, y, 3.05, 0.40, fill=NODE_BG, rounded=True, radius=0.12,
           line=RULE, line_w=0.75)
    glyph(s, 0.44, y + 0.09, 0.22, "message-square", "ink")
    s.text(0.76, y, 2.56, 0.40, text, size=SZ_CAPTION, color=INK, font=BODY_FONT,
           anchor="m")
    group(s, first, f"Correction {i + 1}")
s.arrow(3.47, 2.96, 0.40, 0.14)
m = s.mock(3.97, 1.52, 5.69, 3.40)
m.user("Turn this conversation into a reusable prompt,",
       "so I get the same output every time.")
with m.bubble():
    m.label("AI Assistant")
    m.line("Role: B2B newsletter editor.")
    m.line("Input: {this week's product notes}")
    m.line("Steps: pick 3 updates, lead with the customer benefit.", highlight=True)
    m.line("Format: subject line, 3 short sections, one CTA.", highlight=True)
    m.line("Avoid: jargon, more than 150 words.", highlight=True)
m.fit()
s.caption(3.97, m.y + m.h + 0.16, 5.69,
          "Every correction becomes a line in the prompt.")

# How it works: four stations, then the payoff as a chart.
s = d.how("iterate once, then reuse the recipe")
stations = [("message-square", "grey", "Iterate", "until it's right"),
            ("file-text", "soft", "Capture", "ask for the prompt"),
            ("repeat", "soft", "Reuse", "new chat, new input"),
            ("check", "pink", "Same output", "every run")]
col, bd = 2.44, 0.62
for i, (icon, tone, label, sub) in enumerate(stations):
    cx = 0.34 + i * col + 1.0
    first = mark(s)
    badge(s, cx - bd / 2, 1.55, bd, icon, tone)
    s.text(cx - 1.0, 2.24, 2.0, 0.30, title_case(label), size=14, color=INK,
           bold=True, font=BODY_FONT, align="c")
    s.text(cx - 1.0, 2.54, 2.0, 0.28, title_case(sub), size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT, align="c")
    group(s, first, f"Station {i + 1}")
    if i < len(stations) - 1:
        s.arrow(cx + bd / 2 + 0.14, 1.55 + bd / 2 - 0.06, col - bd - 0.28, 0.12)
# The payoff: one long conversation, then one-turn runs.
s.text(0.34, 3.05, 4.40, 0.28, title_case("Turns needed per run (illustrative)"),
       size=SZ_CAPTION, color=INK, bold=True, font=BODY_FONT)
base, top_h = 4.62, 1.05
for i, turns in enumerate([10, 1, 1, 1, 1]):
    bx = 0.44 + i * 0.84
    bh = top_h * turns / 10
    first = mark(s)
    s.rect(bx, base - bh, 0.52, bh, fill="D8D8DE" if i == 0 else ACCENT)
    s.text(bx - 0.10, base - bh - 0.26, 0.72, 0.24, str(turns), size=SZ_CAPTION,
           color=INK, bold=True, font=BODY_FONT, align="c")
    s.text(bx - 0.14, base + 0.04, 0.80, 0.24, f"Run {i + 1}", size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT, align="c")
    group(s, first, f"Run {i + 1}")
s.rect(0.34, base, 4.30, 0.012, fill=INK)
for i, line in enumerate(["The first run pays for the rest.",
                          "When a run drifts, fix the prompt.",
                          "Test it in a fresh chat first."]):
    first = mark(s)
    glyph(s, 5.10, 3.46 + i * 0.44, 0.24, "check", "pink")
    s.text(5.46, 3.44 + i * 0.44, 4.20, 0.30, line, size=SZ_CAPTION, color=INK,
           font=BODY_FONT, anchor="m")
    group(s, first, f"Takeaway {i + 1}")

# Why it matters: a scorecard, not a grey table.
s = d.why("repeatable tasks stop costing a conversation each")
lx, ax, bx = 0.34, 2.40, 6.10
aw, bw = 3.50, 3.56
s.rect(bx - 0.12, 1.50, bw + 0.24, 3.02, fill=NODE_ACC, rounded=True, radius=0.08)
for x, icon, tone, head in ((ax, "message-square", "grey", "Starting from scratch"),
                            (bx, "bookmark", "pink", "Reusing a meta prompt")):
    first = mark(s)
    badge(s, x, 1.58, 0.44, icon, tone)
    s.text(x + 0.56, 1.58, 3.0, 0.44, title_case(head), size=14,
           color=ACCENT if tone == "pink" else INK, bold=True, font=BODY_FONT,
           anchor="m")
    group(s, first, head)
rows = [("Corrections", ("circle-x", "Re-typed every time"), ("circle-check", "Baked into the prompt")),
        ("Consistency", ("circle-x", "Varies run to run"), ("circle-check", "Same structure each run")),
        ("Sharing", ("circle-x", "Lives in one chat"), ("circle-check", "Paste into a project")),
        ("Failure mode", ("triangle-alert", "Slow and inconsistent"), ("triangle-alert", "Copies a one-off mistake"))]
for i, (label, (ia, ta), (ib, tb)) in enumerate(rows):
    y = 2.22 + i * 0.56
    first = mark(s)
    s.text(lx, y, 1.95, 0.40, title_case(label), size=SZ_CAPTION, color=INK,
           bold=True, font=BODY_FONT, anchor="m")
    glyph(s, ax + 0.08, y + 0.08, 0.24, ia, "ink")
    s.text(ax + 0.44, y, aw - 0.44, 0.40, ta, size=SZ_CAPTION, color=SOFT_INK,
           font=BODY_FONT, anchor="m")
    glyph(s, bx + 0.08, y + 0.08, 0.24, ib, "pink")
    s.text(bx + 0.44, y, bw - 0.44, 0.40, tb, size=SZ_CAPTION, color=INK,
           font=BODY_FONT, anchor="m")
    if i < len(rows) - 1:
        s.rect(lx, y + 0.48, 9.32, 0.008, fill=RULE)
    group(s, first, label)
s.caption(0.34, 4.66, 9.32,
          "Read the generated prompt before you save it.")

d.handoff(
    [
        ("Get one great output", "iterate in the chat"),
        ("Ask for the prompt", "turn this chat into a reusable prompt"),
        ("Run it fresh", "new chat, new input, same result"),
    ],
    lead="Next, we turn one ChatGPT conversation into a prompt we can reuse.",
)

path = d.save(out)
print(f"wrote {path}  ({len(d.prs.slides._sldIdLst)} slides)")
