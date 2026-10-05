#!/usr/bin/env python3
"""Build the five-slide principles deck for working within a context limit."""
import subprocess
import sys
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / ".claude" / "skills" / "principles-deck" / "scripts"))
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from deckkit import (Deck, title_case, ACCENT, INK, SOFT_INK, RULE, NODE_BG,  # noqa: E402
                     NODE_ACC, BODY_FONT, SZ_CAPTION, _rgb)

ASSETS = Path(__file__).resolve().parent / "assets"
LUCIDE = "https://cdn.jsdelivr.net/npm/lucide-static@latest/icons/{}.svg"


def icon(name, color):
    """A Lucide line icon (ISC licence) rendered in one colour, cached in assets/."""
    png = ASSETS / f"icon-{name}-{color}.png"
    if not png.exists():
        ASSETS.mkdir(exist_ok=True)
        svg = urllib.request.urlopen(LUCIDE.format(name)).read().decode()
        tinted = ASSETS / f".{name}-{color}.svg"
        tinted.write_text(svg.replace("currentColor", f"#{color}"))
        subprocess.run(["rsvg-convert", "-w", "256", "-h", "256", str(tinted),
                        "-o", str(png)], check=True)
        tinted.unlink()
    return png


def picture(s, png, x, y, size):
    return s.raw.shapes.add_picture(str(png), Inches(x), Inches(y), Inches(size),
                                    Inches(size))


def badge(s, x, y, d, name, fill=NODE_BG, line=RULE, color=INK, glyph=0.56):
    """An icon inside a filled circle."""
    dot = s.raw.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d),
                                 Inches(d))
    dot.fill.solid()
    dot.fill.fore_color.rgb = _rgb(fill)
    dot.line.color.rgb = _rgb(line)
    dot.line.width = Pt(1.0)
    dot.shadow.inherit = False
    g = d * glyph
    picture(s, icon(name, color), x + (d - g) / 2, y + (d - g) / 2, g)


def group(s, first, name):
    """Group every shape added since index `first`, so it moves as one unit."""
    shapes = s.raw.shapes
    grp = shapes.add_group_shape(list(shapes)[first:])
    grp.name = name
    return grp


def label(s, x, y, w, text, size=SZ_CAPTION, color=INK, bold=False, align="c",
          h=0.26):
    return s.text(x, y, w, h, text, size=size, color=color, bold=bold,
                  font=BODY_FONT, align=align, anchor="m")


out = Path(sys.argv[1] if len(sys.argv) > 1 else
           Path(__file__).resolve().parent / "overcoming-token-limit.pptx")
d = Deck()

d.title(
    "Overcoming the Token Limit",
    "How to keep the right information in context when the whole task will not fit at once.",
)

# 2. The context window as a capacity bar: one that fits, one that overflows.
s = d.what("the context window is a finite working surface")
X0, LIMIT = 0.34, 7.60                      # bar start, window width
SEGS = {  # name: (icon, fill, text colour)
    "Instructions": ("scroll-text", "E4E5EA", INK),
    "Conversation": ("messages-square", "ECECF0", INK),
    "Source": ("file-text", "F5F5F7", INK),
    "Answer Space": ("pen-line", NODE_ACC, ACCENT),
}


def bar(y, parts, title, title_icon, title_color):
    first = len(s.raw.shapes)
    badge(s, X0, y, 0.30, title_icon, fill="FFFFFF", line="FFFFFF",
          color=title_color, glyph=0.9)
    label(s, X0 + 0.38, y + 0.02, 4.0, title_case(title), size=14, bold=True,
          color=title_color, align="l")
    by, bh, x = y + 0.40, 0.56, X0
    for name, frac, over in parts:
        ic, fill, ink = SEGS[name]
        w = LIMIT * frac
        seg = s.rect(x, by, w, bh, fill=None if over else fill,
                     line=ACCENT if name == "Answer Space" else RULE, line_w=1.0)
        if over:
            seg.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        picture(s, icon(ic, ACCENT if name == "Answer Space" else "424242"),
                x + w / 2 - 0.13, by + 0.15, 0.26)
        label(s, x, by + bh + 0.04, w, "Cut Off" if over else name,
              color=ACCENT if (over or name == "Answer Space") else SOFT_INK,
              bold=over)
        x += w
    group(s, first, title)


bar(1.52, [("Instructions", 0.17, False), ("Conversation", 0.20, False),
           ("Source", 0.33, False), ("Answer Space", 0.30, False)],
    "Fits", "circle-check", "424242")
bar(3.02, [("Instructions", 0.17, False), ("Conversation", 0.20, False),
           ("Source", 0.63, False), ("Answer Space", 0.20, True)],
    "Too much source", "triangle-alert", ACCENT)
# The window limit: one line through both bars.
s.rect(X0 + LIMIT - 0.01, 1.80, 0.025, 2.82, fill=ACCENT)
label(s, X0 + LIMIT - 1.00, 1.52, 2.00, title_case("Window limit"), color=ACCENT,
      bold=True)
s.caption(0.34, 4.70, 9.32,
          "Everything shares one window. Too much source leaves no room to answer.")

# 3. The staged workflow as a pipeline: split, extract, store, synthesise.
s = d.how("a long task becomes a sequence of bounded passes")
cols = [(0.34, 1.30, "Split"), (2.04, 1.80, "Extract"), (4.30, 2.70, "Store"),
        (7.46, 2.20, "Synthesise")]
for i, (cx, cw, name) in enumerate(cols, 1):
    first = len(s.raw.shapes)
    s.disc(cx + cw / 2 - 0.62, 1.54, 0.30, str(i), size=12)
    label(s, cx + cw / 2 - 0.26, 1.56, 1.40, name, size=14, bold=True, align="l")
    group(s, first, f"Step {i}")
row_y, row_h, gap = 2.08, 0.52, 0.12
rows = [row_y + k * (row_h + gap) for k in range(4)]
mid = rows[0] + (rows[-1] + row_h - rows[0]) / 2
# The long source.
first = len(s.raw.shapes)
picture(s, icon("file-text", "424242"), 0.34 + 0.65 - 0.45, mid - 0.62, 0.90)
label(s, 0.34, mid + 0.34, 1.30, title_case("Long source"), color=SOFT_INK)
group(s, first, "Long source")
s.arrow(1.60, mid - 0.06, 0.36, 0.12)
for k, y in enumerate(rows, 1):
    first = len(s.raw.shapes)
    s.rect(2.04, y, 1.80, row_h, fill=NODE_BG, rounded=True, radius=0.05,
           line=RULE, line_w=0.75)
    label(s, 2.04, y, 1.80, f"Section {k}", h=row_h)
    s.arrow(3.90, y + row_h / 2 - 0.05, 0.34, 0.10)
    s.rect(4.30, y, 2.70, row_h, fill="FFFFFF", rounded=True, radius=0.05,
           line=RULE, line_w=0.75)
    picture(s, icon("list-checks", "424242"), 4.40, y + 0.13, 0.26)
    label(s, 4.74, y, 1.40, "Same Fields", h=row_h, align="l")
    s.rect(6.26, y + 0.12, 0.60, 0.28, fill=NODE_ACC, rounded=True, radius=0.14,
           line=ACCENT, line_w=0.75)
    label(s, 6.26, y + 0.12, 0.60, f"§{k}", color=ACCENT, bold=True, h=0.28)
    group(s, first, f"Section {k}")
# Converge the summaries into one checked answer.
s.rect(7.10, rows[0] + row_h / 2, 0.02, rows[-1] - rows[0], fill=RULE)
for y in rows:
    s.rect(7.00, y + row_h / 2 - 0.01, 0.12, 0.02, fill=RULE)
s.arrow(7.12, mid - 0.06, 0.30, 0.12)
first = len(s.raw.shapes)
s.node(7.46, mid - 0.62, 2.20, 1.24, "Final answer", sub="checked against sources",
       accent=True)
group(s, first, "Final answer")
s.caption(0.34, 4.66, 9.32,
          "One summary shape for every section. References let you reopen the source.")

# 4. Why it matters: one giant chat against a staged workflow, row by row.
s = d.why("staging protects the signal you need later")
LX, AX, AW, BX, BW = 0.34, 3.00, 3.10, 6.30, 3.36
s.rect(BX - 0.12, 1.50, BW + 0.12, 3.10, fill=NODE_ACC, rounded=True, radius=0.06)
for x, name, ic, color in ((AX, "One giant chat", "message-square", "424242"),
                           (BX, "Staged workflow", "layers", ACCENT)):
    first = len(s.raw.shapes)
    picture(s, icon(ic, color), x, 1.62, 0.30)
    label(s, x + 0.40, 1.62, 2.80, title_case(name), size=14, bold=True,
          color=color, align="l", h=0.30)
    group(s, first, name)
table = [
    ("filter", "Selection", "Relevant and irrelevant mixed", "Each pass has one purpose"),
    ("clipboard-list", "Intermediate work", "Buried in the conversation",
     "Captured in a fixed shape"),
    ("search-check", "Verification", "Claims are hard to trace",
     "Back to a named section"),
    ("triangle-alert", "Failure mode", "Truncation or missed detail",
     "A summary drops a nuance"),
]
y = 2.08
for ic, name, one, staged in table:
    first = len(s.raw.shapes)
    badge(s, LX, y + 0.06, 0.44, ic)
    label(s, LX + 0.56, y, 2.00, title_case(name), bold=True, align="l", h=0.56)
    label(s, AX, y, AW, one, color=SOFT_INK, align="l", h=0.56)
    label(s, BX, y, BW - 0.12, staged, color=INK, align="l", h=0.56)
    s.rect(LX, y + 0.60, 9.32, 0.008, fill=RULE)
    group(s, first, name)
    y += 0.62

d.handoff(
    [
        ("Choose a long source", "find its sections"),
        ("Define the fields", "use one summary shape"),
        ("Process in passes", "keep source references"),
        ("Synthesise and check", "verify key claims"),
    ],
    lead="Next, we turn one oversized task into a small, repeatable sequence.",
)

path = d.save(out)
print(f"wrote {path}  ({len(d.prs.slides._sldIdLst)} slides)")
