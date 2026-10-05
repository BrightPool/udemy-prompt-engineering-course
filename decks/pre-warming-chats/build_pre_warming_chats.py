#!/usr/bin/env python3
"""Build the five-slide principles deck for pre-warming a conversation."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / ".claude" / "skills" / "principles-deck" / "scripts"))
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from deckkit import (Deck, title_case, _rgb, ACCENT, INK, SOFT_INK, RULE,  # noqa: E402
                     NODE_BG, NODE_ACC, BODY_FONT, SZ_CAPTION, SZ_DIAGRAM)

ASSETS = Path(__file__).resolve().parent / "assets"   # Lucide icons, ISC licence


def mark(s):
    """Remember where the shape list ends, so the next shapes can be grouped."""
    return len(s.raw.shapes)


def group(s, start, name):
    """Group every shape added since `start`, so it moves as one unit."""
    shapes = s.raw.shapes
    g = shapes.add_group_shape(list(shapes)[start:])
    g.name = name
    return g


def badge(s, x, y, d, icon, fill=NODE_ACC, line=ACCENT, glyph=0.55):
    """An icon in a circle."""
    dot = s.raw.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                                 Inches(d), Inches(d))
    dot.fill.solid()
    dot.fill.fore_color.rgb = _rgb(fill)
    dot.line.color.rgb = _rgb(line)
    dot.line.width = Pt(1.0)
    dot.shadow.inherit = False
    g = d * glyph
    s.raw.shapes.add_picture(str(ASSETS / f"{icon}.png"), Inches(x + (d - g) / 2),
                             Inches(y + (d - g) / 2), Inches(g), Inches(g))


def pill(s, x, y, w, label, icon=None, accent=False, h=0.46):
    """A timeline step: rounded pill, optional icon on the left."""
    s.rect(x, y, w, h, fill=NODE_ACC if accent else NODE_BG, rounded=True,
           radius=0.10, line=ACCENT if accent else RULE, line_w=1.0 if accent else 0.75)
    tx = x + 0.08
    if icon:
        g = 0.22
        s.raw.shapes.add_picture(str(ASSETS / f"{icon}.png"), Inches(x + 0.12),
                                 Inches(y + (h - g) / 2), Inches(g), Inches(g))
        tx = x + 0.38
    s.text(tx, y, x + w - 0.08 - tx, h, title_case(label), size=SZ_CAPTION,
           color=ACCENT if accent else INK, bold=accent, font=BODY_FONT,
           align="c", anchor="m")


out = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parent / "pre-warming-chats.pptx")
d = Deck()

d.title(
    "Pre-Warming Chats",
    "How a short setup exchange gives later requests a shared working context.",
)

# What it is: the briefing, with its four parts called out beside it.
s = d.what("the conversation starts with a compact briefing")
m = s.mock(0.40, 1.52, 4.60, 3.58, generic=True)
m.user("We are launching a reporting feature.")
m.line("Audience: existing customers.")
m.line("Source: the attached release notes.")
m.line("Launch Friday. Keep emails under 150 words.")
with m.bubble():
    m.assistant("Got it: reporting launch for existing",
                "customers, from the release notes, by Friday.")
m.divider()
m.user("Now draft the announcement email.", highlight=True)
m.fit()

rx, rw = 5.45, 4.21
s.text(rx, 1.56, rw, 0.30, title_case("A good brief has four parts"), size=14,
       color=INK, bold=True, font=BODY_FONT)
parts = [
    ("target-E91D63", "Goal", "A reporting feature launch"),
    ("users-E91D63", "Audience", "Existing customers"),
    ("file-text-E91D63", "Source", "The release notes"),
    ("ruler-E91D63", "Constraints", "Friday, under 150 words"),
]
py = 2.00
for icon, label, example in parts:
    start = mark(s)
    badge(s, rx, py, 0.44, icon)
    s.text(rx + 0.58, py - 0.02, rw - 0.58, 0.26, title_case(label), size=SZ_DIAGRAM,
           color=INK, bold=True, font=BODY_FONT)
    s.text(rx + 0.58, py + 0.22, rw - 0.58, 0.24, example, size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT)
    group(s, start, f"Brief part: {label}")
    py += 0.62
s.text(rx, py + 0.06, rw, 0.30, "Brief once. Later turns build on it.",
       size=SZ_CAPTION, color=ACCENT, bold=True, font=BODY_FONT)

# How it works: every turn, the model reads the whole stack.
s = d.how("the current request sits on top of earlier context")
stack = [  # top to bottom: newest first
    ("message-square-FFFFFF", "Latest request", "This Turn's Task", True),
    ("file-text-424242", "Source material", "Facts It May Use", False),
    ("users-424242", "Audience", "Who It Serves", False),
    ("target-424242", "Stable goal", "What the Work Is For", False),
]
sx, sw, bh, gap = 0.34, 4.90, 0.60, 0.08
by = 1.56
for icon, label, note, accent in stack:
    start = mark(s)
    s.rect(sx, by, sw, bh, fill=NODE_ACC if accent else NODE_BG, rounded=True,
           radius=0.05, line=ACCENT if accent else RULE, line_w=1.0 if accent else 0.75)
    if accent:
        badge(s, sx + 0.10, by + 0.09, 0.42, icon, fill=ACCENT)
    else:
        badge(s, sx + 0.10, by + 0.09, 0.42, icon, fill="FFFFFF", line=RULE)
    s.text(sx + 0.66, by, 2.10, bh, title_case(label), size=SZ_DIAGRAM,
           color=ACCENT if accent else INK, bold=True, font=BODY_FONT, anchor="m")
    s.text(sx + 2.70, by, sw - 2.86, bh, note, size=SZ_CAPTION, color=SOFT_INK,
           font=BODY_FONT, align="r", anchor="m")
    group(s, start, f"Stack: {title_case(label)}")
    by += bh + gap
top, bottom = 1.56, by - gap
# Label the older layers as the brief from turn one.
s.text(sx, bottom + 0.06, sw, 0.26, "Turn 1: the brief. Turn 5: the latest request.",
       size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT, align="c")

# Bracket: the whole stack is read together, every turn.
start = mark(s)
bx = sx + sw + 0.22
s.rect(bx, top, 0.02, bottom - top, fill=ACCENT)
s.rect(bx - 0.10, top, 0.12, 0.02, fill=ACCENT)
s.rect(bx - 0.10, bottom - 0.02, 0.12, 0.02, fill=ACCENT)
s.arrow(bx + 0.10, (top + bottom) / 2 - 0.06, 0.40, 0.12)
group(s, start, "Bracket")
tx = bx + 0.66
s.text(tx, 2.10, 9.66 - tx, 0.64, title_case("The model reads the whole stack, every turn"),
       size=14, color=ACCENT, bold=True, font=BODY_FONT, spacing=1.15)
s.text(tx, 2.86, 9.66 - tx, 0.50,
       "It is not smarter. It just has a clearer history.",
       size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT, spacing=1.15)

# Tip strip: correct the brief before the real work starts.
start = mark(s)
badge(s, 0.34, 4.64, 0.40, "circle-check-E91D63")
s.text(0.86, 4.64, 8.80, 0.40,
       "Ask it to restate the brief, and fix errors before the real task.",
       size=SZ_CAPTION, color=INK, font=BODY_FONT, anchor="m")
group(s, start, "Tip: restate the brief")

# Why it matters: two timelines, cold versus prepared.
s = d.why("shared context improves continuity across turns")


def timeline(s, y, heading, icon, steps, band=None):
    """A labelled row of pills joined by arrows. `band` = (first, last, label)."""
    start = mark(s)
    badge(s, 0.34, y, 0.34, icon, fill="FFFFFF", line=RULE)
    s.text(0.78, y, 4.00, 0.34, title_case(heading), size=14, color=INK, bold=True,
           font=BODY_FONT, anchor="m")
    n, aw = len(steps), 0.30
    w = (9.32 - aw * (n - 1)) / n
    py = y + 0.48
    xs = []
    for i, (label, step_icon, accent) in enumerate(steps):
        x = 0.34 + i * (w + aw)
        xs.append(x)
        pill(s, x, py, w, label, icon=step_icon, accent=accent)
        if i < n - 1:
            s.arrow(x + w + 0.06, py + 0.17, aw - 0.12, 0.12)
    if band:
        a, b, text = band
        x0, x1 = xs[a], xs[b] + w
        s.rect(x0, py + 0.58, x1 - x0, 0.03, fill=ACCENT)
        s.text(x0, py + 0.64, x1 - x0, 0.26, title_case(text), size=SZ_CAPTION,
               color=ACCENT, bold=True, font=BODY_FONT, align="c")
    group(s, start, f"Timeline: {title_case(heading)}")


timeline(s, 1.52, "Cold request", "zap-424242", [
    ("One giant prompt", None, False),
    ("First output", None, False),
    ("Fix afterwards", "triangle-alert-E91D63", False),
    ("Start over", "refresh-ccw-424242", False),
])
timeline(s, 2.78, "Prepared conversation", "message-square-E91D63", [
    ("Brief", None, False),
    ("Confirm", "circle-check-E91D63", True),
    ("Task 1", None, False),
    ("Task 2", None, False),
    ("Task 3", None, False),
], band=(2, 4, "One shared brief, reused"))

start = mark(s)
badge(s, 0.34, 4.44, 0.40, "triangle-alert-E91D63")
s.text(0.86, 4.44, 8.80, 0.40,
       "Watch for stale or bloated history. When the brief changes, open a fresh chat.",
       size=SZ_CAPTION, color=INK, font=BODY_FONT, anchor="m")
group(s, start, "Failure mode")

d.handoff(
    [
        ("Open a fresh chat", "start clean"),
        ("Give a compact brief", "goal and audience"),
        ("Confirm the context", "correct the summary"),
        ("Give the real task", "reuse the shared brief"),
    ],
    lead="Next, we prepare the chat before asking for the deliverable.",
)

path = d.save(out)
print(f"wrote {path}  ({len(d.prs.slides._sldIdLst)} slides)")
