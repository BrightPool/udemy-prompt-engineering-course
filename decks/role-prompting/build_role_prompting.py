#!/usr/bin/env python3
"""Build the five-slide principles deck for role prompting."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / ".claude" / "skills" / "principles-deck" / "scripts"))
from deckkit import Deck, title_case  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from deckkit import (ACCENT, ARROW, INK, SOFT_INK, RULE, NODE_BG, NODE_ACC,  # noqa: E402
                     BODY_FONT, HEAD_FONT, SZ_CAPTION, _rgb)

# Visual helpers for this deck (icons are Lucide, ISC licence, rendered to assets/).
ASSETS = Path(__file__).resolve().parent / "assets"
TONES = {"pink": (ACCENT, ACCENT, "white"), "soft": (NODE_ACC, ACCENT, "pink"),
         "grey": ("ECECF0", RULE, "ink")}


def mark(s):
    """Shape count now; pass it to group() to bundle everything drawn since."""
    return len(s.raw.shapes)


def group(s, first, name=None):
    """Group every shape drawn since `first`, so it moves as one unit."""
    shapes = s.raw.shapes
    g = shapes.add_group_shape(list(shapes)[first:])
    if name:
        g.name = name
    return g


def glyph(s, x, y, size, icon, colour="pink"):
    """A bare icon."""
    s.raw.shapes.add_picture(str(ASSETS / f"i-{icon}-{colour}.png"),
                             Inches(x), Inches(y), Inches(size), Inches(size))


def oval(s, x, y, d, fill, line=None):
    c = s.raw.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d),
                               Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = _rgb(fill)
    if line:
        c.line.color.rgb = _rgb(line)
        c.line.width = Pt(1.0)
    else:
        c.line.fill.background()
    c.shadow.inherit = False
    return c


def badge(s, x, y, d, icon, tone="pink"):
    """A round icon badge: pink (solid), soft (pink tint) or grey."""
    fill, line, colour = TONES[tone]
    oval(s, x, y, d, fill, line)
    g = d * 0.54
    glyph(s, x + (d - g) / 2, y + (d - g) / 2, g, icon, colour)

out = Path(sys.argv[1] if len(sys.argv) > 1 else
           Path(__file__).resolve().parent / "role-prompting.pptx")
d = Deck()

d.title(
    "Role Prompting",
    "How telling the model who it is changes the tone, focus and framing of its answer.",
)

# What it is: one request, three roles. Only the role line changes.
s = d.what("a role sets the perspective of the answer")
mw, gap, top = 2.96, 0.22, 2.02
panels = []
for i, (role, icon, reply) in enumerate([
    ("a copywriter", "pen-tool",
     ["Lead with the benefit:", "\"Sign in with one click,", "from Friday.\""]),
    ("a senior engineer", "code",
     ["Avoid a Friday release.", "Ship Tuesday behind a", "feature flag."]),
    ("a product manager", "clipboard-list",
     ["Define success first:", "fewer failed logins in", "the first week."]),
]):
    x = 0.34 + i * (mw + gap)
    first = mark(s)
    badge(s, x, 1.50, 0.40, icon, "pink")
    s.text(x + 0.52, 1.50, mw - 0.52, 0.40, title_case(role.split(" ", 1)[1]),
           size=14, color=INK, bold=True, font=BODY_FONT, anchor="m")
    group(s, first, f"Role {i + 1}")
    m = s.mock(x, top, mw, 3.00, app="AI Assistant", tag="Example")
    m.user(f"Act as {role}.", "Review: Friday login launch.")
    with m.bubble():
        m.label("AI Assistant")
        for ln in reply:
            m.line(ln, highlight=True)
    m.fit()
    panels.append(m)
# Match the windows so the comparison reads as a set.
tallest = max(p.h for p in panels)
for p in panels:
    p.h = tallest
    p._panel.height = int(tallest * 914400)
s.caption(0.34, top + tallest + 0.14, 9.32,
          "Same plan. Only the role changed, and with it the priorities.")

# How it works: what a role moves, what it leaves alone, and the evidence.
s = d.how("a role shifts emphasis, not knowledge")
cols = [
    (0.34, 3.80, "What a role changes", "soft", ACCENT,
     [("mic", "Vocabulary and tone"), ("list-ordered", "Which details come first"),
      ("target", "What a good answer looks like"),
      ("help-circle", "The questions it asks back")]),
    (4.28, 3.30, "What it doesn't change", "grey", INK,
     [("brain", "What the model knows"), ("shield-check", "Whether facts are right"),
      ("circle-check", "Your need to check it")]),
]
for x, w, head, tone, colour, rows in cols:
    s.text(x, 1.52, w, 0.34, title_case(head), size=14, color=colour, bold=True,
           font=BODY_FONT)
    for j, (icon, label) in enumerate(rows):
        y = 1.98 + j * 0.60
        first = mark(s)
        badge(s, x, y, 0.44, icon, tone)
        s.text(x + 0.58, y, w - 0.58, 0.44, title_case(label), size=SZ_CAPTION,
               color=INK, font=BODY_FONT, anchor="m")
        group(s, first, label)
# The evidence, as a stat card.
cx, cw = 7.74, 1.92
first = mark(s)
s.rect(cx, 1.52, cw, 2.94, fill=NODE_ACC, rounded=True, radius=0.08,
       line=ACCENT, line_w=0.75)
for y, big, small in ((1.62, "162", "Roles Tested"), (2.48, "2,410", "Factual Questions")):
    s.text(cx, y, cw, 0.50, big, size=26, color=ACCENT, bold=True, font=HEAD_FONT,
           align="c")
    s.text(cx, y + 0.50, cw, 0.26, small, size=SZ_CAPTION, color=INK,
           font=BODY_FONT, align="c")
s.rect(cx + 0.25, 3.44, cw - 0.50, 0.008, fill=ACCENT)
s.text(cx + 0.12, 3.54, cw - 0.24, 0.80, "No Average Accuracy Gain", size=SZ_CAPTION,
       color=ACCENT, bold=True, font=BODY_FONT, align="c", spacing=1.15)
group(s, first, "Study")
s.caption(0.34, 4.62, 9.32, "Source: persona prompting study, EMNLP 2024.")

# Why it matters: a spectrum from vague to specific, plus the failure mode.
s = d.why("specific roles beat vague ones")
ax0, ax1, ay = 0.60, 9.30, 2.80
s.rect(ax0, ay - 0.03, ax1 - ax0 - 0.20, 0.06, fill=RULE, rounded=True)
s.arrow(ax1 - 0.26, ay - 0.08, 0.26, 0.16, color="B8B8C0")
stops = [
    (1.70, 0.24, "E0E0E6", "\"Act as an expert.\"", "Barely changes the answer"),
    (5.00, 0.32, NODE_ACC, "\"Act as a senior engineer.\"", "Right vocabulary"),
    (8.10, 0.40, ACCENT, "\"Act as a senior engineer reviewing a junior's code for security.\"",
     "Right vocabulary, priorities and checks"),
]
for cxs, dd, fill, quote, result in stops:
    first = mark(s)
    oval(s, cxs - dd / 2, ay - dd / 2, dd, fill, ACCENT if fill != "E0E0E6" else RULE)
    s.text(cxs - 1.40, 1.56, 2.80, 0.90, quote, size=SZ_CAPTION, color=INK,
           font=BODY_FONT, align="c", anchor="b", spacing=1.15)
    s.text(cxs - 1.40, ay + 0.34, 2.80, 0.50, title_case(result), size=SZ_CAPTION,
           color=ACCENT if fill == ACCENT else SOFT_INK, bold=fill == ACCENT,
           font=BODY_FONT, align="c", spacing=1.15)
    group(s, first, quote)
first = mark(s)
s.rect(0.34, 4.00, 9.32, 0.62, fill=NODE_ACC, rounded=True, radius=0.10,
       line=ACCENT, line_w=0.75)
glyph(s, 0.56, 4.15, 0.32, "triangle-alert", "pink")
s.text(1.04, 4.00, 8.50, 0.62,
       "Failure mode: a convincing voice can hide a wrong fact. Still check it.",
       size=SZ_CAPTION, color=INK, bold=True, font=BODY_FONT, anchor="m")
group(s, first, "Failure mode")

d.handoff(
    [
        ("Ask with no role", "see the default answer"),
        ("Add a role", "act as a copywriter"),
        ("Swap the role", "act as a senior engineer"),
    ],
    lead="Next, we send the same request in ChatGPT with different roles.",
)

path = d.save(out)
print(f"wrote {path}  ({len(d.prs.slides._sldIdLst)} slides)")
