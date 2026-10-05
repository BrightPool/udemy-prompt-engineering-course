#!/usr/bin/env python3
"""
"What are Tokens?" - the principles half of the re-film.

Run:
    uv run --with python-pptx python decks/tokens/build_tokens.py decks/tokens/tokens.pptx
    uv run --with python-pptx python .claude/skills/principles-deck/scripts/check_deck.py \
        decks/tokens/tokens.pptx

Arc: a token is a chunk of characters, not a word (what) - text is split,
numbered and predicted back one token at a time inside one shared window (how) -
and every token is billed in one of three buckets (why).

No prices anywhere in here. The three buckets are durable; the numbers are not.
Today's rates are in NOTES.md for James to say aloud.
"""
import sys
from pathlib import Path
from types import SimpleNamespace

SKILL = Path(__file__).resolve().parents[2] / ".claude" / "skills" / "principles-deck"
sys.path.insert(0, str(SKILL / "scripts"))
from deckkit import (  # noqa: E402
    ACCENT, BODY_FONT, INK, MUTED, NODE_ACC, NODE_BG, PANEL, ROLE_AI, ROLE_USER,
    RULE, SOFT_INK, SZ_CAPTION, SZ_DIAGRAM, TEXT, Deck,
)
from code_image import code_image  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

ASSETS = Path(__file__).resolve().parent / "assets"

out = Path(sys.argv[1] if len(sys.argv) > 1 else "decks/tokens/tokens.pptx")


# ---------------------------------------------------------------- helpers
def picture(s, path, x, y, w, h=None, ratio=None):
    """Place an image. Height derives from the real aspect ratio unless given."""
    if h is None:
        h = w / ratio
    s.raw.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    return y + h


def code(s, x, y, w, lines, title="Python", size=11.0, lead=0.235):
    """A code panel. Source Code Pro is monospaced, so it renders as real code."""
    h = 0.34 + len(lines) * lead + 0.20
    s.rect(x, y, w, h, fill=PANEL, rounded=True, radius=0.06)
    s.rect(x, y, w, 0.30, fill="25262C", rounded=True, radius=0.06)
    s.rect(x, y + 0.19, w, 0.11, fill="25262C")
    s.text(x + 0.16, y + 0.06, w - 0.32, 0.20, title,
           size=9.0, color=MUTED, font=BODY_FONT, bold=True)
    for i, ln in enumerate(lines):
        s.text(x + 0.20, y + 0.40 + i * lead, w - 0.40, lead, ln,
               size=size, color=TEXT, font=BODY_FONT, wrap=False)
    return y + h


def prefix_bar(s, x, y, w, label, note, cached_frac, h=0.62,
               cached_text="matched - cheap", fresh_text="new - full price",
               broken=False):
    """One prompt drawn as a bar: the part the provider recognised, then the rest.

    Caching matches from the FIRST character, so what earns the discount is a
    prefix, not a set of blocks. A bar shows that; stacked boxes do not.
    """
    # label and its note stack in the left column, beside the bar - not beneath
    # it, which pushed each row 0.4" taller than it needed to be.
    s.text(x, y + 0.02, 1.66, 0.26, label, size=SZ_CAPTION, color=INK,
           bold=True, font=BODY_FONT)
    s.text(x, y + 0.28, 1.66, 0.40, note, size=SZ_CAPTION, color=SOFT_INK,
           font=BODY_FONT, spacing=1.15)
    bx, bw = x + 1.80, w - 1.80
    cw = bw * cached_frac
    if cw > 0.04:
        s.rect(bx, y, cw - 0.02, h, fill=NODE_ACC, rounded=True, radius=0.05,
               line=ACCENT, line_w=1.0)
        if cw > 1.30:
            s.text(bx + 0.10, y, cw - 0.24, h, cached_text, size=SZ_CAPTION,
                   color=INK, font=BODY_FONT, anchor="m", align="c")
    if bw - cw > 0.04:
        s.rect(bx + cw, y, bw - cw, h, fill=NODE_BG, rounded=True, radius=0.05,
               line=ACCENT if broken else RULE, line_w=1.0 if broken else 0.75)
        if bw - cw > 1.30:
            s.text(bx + cw + 0.10, y, bw - cw - 0.20, h, fresh_text,
                   size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT,
                   anchor="m", align="c")
    return y + h + 0.16


def budget_bar(s, x, y, w, segments, h=0.72):
    """One fixed bar cut into the things competing for it.

    Layers stack; a bar divides. The point of a context window is that the total
    is fixed and every part is spending the same budget - which a divided bar
    shows and a stack cannot.
    """
    total = sum(seg[2] for seg in segments)
    s.text(x, y - 0.34, w, 0.26, "The context window - one fixed total",
           size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT, align="c")
    s.rect(x, y - 0.10, 0.012, 0.10, fill=RULE)
    s.rect(x + w - 0.012, y - 0.10, 0.012, 0.10, fill=RULE)
    s.rect(x, y - 0.10, w, 0.012, fill=RULE)
    cx = x
    for label, sub, weight, accent in segments:
        sw = w * weight / total
        s.rect(cx, y, sw - 0.02, h, fill=NODE_ACC if accent else NODE_BG,
               line=ACCENT if accent else RULE, line_w=1.0 if accent else 0.75)
        s.text(cx + 0.06, y, sw - 0.14, h, label, size=SZ_CAPTION, color=INK,
               bold=accent, font=BODY_FONT, align="c", anchor="m")
        s.text(cx + 0.06, y + h + 0.10, sw - 0.14, 0.46, sub,
               size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT, align="c",
               spacing=1.15)
        cx += sw
    return y + h + 0.60


def _mono_width(text, size_pt):
    """Rendered width of Source Code Pro, which is monospaced at 0.6 em."""
    return len(text) * size_pt * 0.60 / 72.0


def chips(s, x, y, w, tokens, accent=(), size=SZ_DIAGRAM, h=0.40,
          gap=0.09, line_gap=0.11):
    """A row of token pills - the same sentence, cut where the tokeniser cuts.

    Wraps to the next line when it runs out of width. Returns the bottom y.
    """
    cur_x, cur_y = x, y
    for i, tok in enumerate(tokens):
        cw = _mono_width(tok, size) + 0.28
        if cur_x > x and cur_x + cw > x + w:
            cur_x, cur_y = x, cur_y + h + line_gap
        hot = i in accent
        s.rect(cur_x, cur_y, cw, h, fill=NODE_ACC if hot else NODE_BG,
               rounded=True, radius=0.05,
               line=ACCENT if hot else RULE, line_w=1.0 if hot else 0.75)
        s.text(cur_x, cur_y, cw, h, tok, size=size, color=INK, bold=hot,
               font=BODY_FONT, align="c", anchor="m")
        cur_x += cw + gap
    return cur_y + h


def bar_row(s, y, label, note, bar_w, rate, accent=False, h=0.72):
    """A labelled bar carrying its rate. Width is the ordering, not a scale."""
    s.text(0.34, y, 2.50, 0.32, label, size=SZ_DIAGRAM, color=INK, bold=True,
           font=BODY_FONT, anchor="b")
    s.text(0.34, y + 0.36, 2.50, 0.32, note, size=SZ_CAPTION, color=SOFT_INK,
           font=BODY_FONT, anchor="t")
    s.rect(3.00, y + 0.14, bar_w, 0.44, fill=NODE_ACC if accent else NODE_BG,
           rounded=True, radius=0.05,
           line=ACCENT if accent else RULE, line_w=1.25 if accent else 0.75)
    s.text(3.16, y + 0.14, bar_w - 0.32, 0.44, rate, size=SZ_CAPTION,
           color=INK, font=BODY_FONT, anchor="m")
    return y + h


d = Deck()

# ---------------------------------------------------------------- open
d.title("What are Tokens?",
        "The unit of text a model reads, writes, counts and charges for.")

# ============================================================== WHAT IS IT?
# Leads with the answer, like slide 2 of every other deck. The title slide has
# already asked the question.
s = d.what("a sequence of characters, not a word")
s.node(0.34, 1.56, 5.40, 0.56, "Unbelievably, tokenisation is unpredictable.")
s.arrow(2.88, 2.22, 0.28, 0.22, direction="down")
bottom = chips(s, 0.34, 2.58, 5.40,
               ["Un", "bel", "ievably", ",", "token", "isation", "is",
                "unpredictable", "."],
               accent=(0, 1, 2))
s.caption(0.34, bottom + 0.16, 5.40, "Four words. Nine tokens.")
panel = SimpleNamespace(x=0.34, y=1.56, w=5.40, h=bottom - 1.56)
s.beside(panel,
         "A token is a sequence of characters common enough to be stored as "
         "one unit. Frequent words survive whole; rarer ones arrive in pieces.")

s = d.content("Why the count is never the word count")
s.table(0.34, 1.56, 9.32, 2.10, [
    ["What you send", "Characters", "Tokens", "Chars per token"],
    ["A short English sentence", "37", "7", "5.3"],
    ["The same sentence in Japanese", "20", "16", "1.3"],
    ["One line of Python", "64", "14", "4.6"],
    ["A written-out money figure", "43", "14", "3.1"],
], col_widths=[3.9, 1.7, 1.5, 2.2])
s.caption(0.34, 3.86, 9.32,
          "English runs at roughly four characters to the token. Other scripts, "
          "code and long numbers all break into smaller pieces, so the identical "
          "sentence can cost twice as many tokens in Japanese as in English.")

s = d.exercise("See a sentence become tokens", minutes=1)
_r1 = picture(s, ASSETS / "tokenizer-text.png", 0.34, 1.14, 3.46, ratio=691 / 535)
picture(s, ASSETS / "tokenizer-ids.png", 3.98, 1.14, 3.46, ratio=699 / 553)
s.text(0.34, _r1 + 0.12, 3.46, 0.26, "The same seven tokens, coloured",
       size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT, align="c")
s.text(3.98, _r1 + 0.12, 3.46, 0.26, "and the numbers behind them",
       size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT, align="c")
s.steps(7.62, 1.14, 2.04, [
    "Open the OpenAI tokeniser.",
    "Paste a sentence of your own.",
    "Switch to Token IDs.",
    "Try it in another language.",
], h=0.62, gap=0.10, size=SZ_CAPTION)
_link = s.caption(7.62, 4.24, 2.04, "platform.openai.com/tokenizer",
                  color=ACCENT)
# make it a real, clickable link so it works from the deck itself
for _para in _link.text_frame.paragraphs:
    for _r in _para.runs:
        _r.hyperlink.address = "https://platform.openai.com/tokenizer"

s = d.content("What a token is not")
s.compare(1.56,
          "Not this", ["One token, one word",
                       "The same count in every language",
                       "Something you can eyeball"],
          "This", ["A frequent sequence of characters",
                   "A fixed vocabulary, learned once",
                   "Countable, before you send it"],
          h=2.00)
s.caption(0.34, 3.72, 9.32,
          "Every model ships with its vocabulary already fixed, so a given "
          "sentence has one token count and you can know it in advance.")

# ============================================================== HOW IT WORKS
s = d.how("split, number, predict, repeat")
s.flow(0.34, 1.56, 9.32,
       [("Your text", "typed as normal"),
        ("Split", "into known chunks"),
        ("Numbered", "one ID per chunk"),
        ("Predicted", "the next ID, then the next")],
       h=1.20, accent_last=True)
s.caption(0.34, 2.96, 9.32,
          "The model never sees your letters. It sees a list of numbers, and it "
          "answers by choosing the next number over and over - which is why a "
          "reply arrives in pieces rather than all at once:")
chips(s, 0.34, 3.76, 9.32,
      ["Rough", "ly", "four", "characters", "to", "the", "token", "…"],
      accent=(7,))

# The context window, shown as the thing it actually is: a frame drawn round a
# real conversation. The mock is a normal ChatGPT thread; the pink rule down its
# side is what the model re-reads this turn, and the accented row shows how small
# the part you just typed really is.
s = d.content("What has to fit in the context window")

MX, MY, MW = 0.40, 1.56, 6.04
m = s.mock(MX, MY, MW, 3.36)
m.label("System").line("You are a helpful assistant.")
m.gap(0.05)
m.label(ROLE_USER).line("Summarise this report in five points.")
m.gap(0.05)
with m.bubble():
    m.label(ROLE_AI).line("1. Revenue is flat.  2. Churn is up.  3. \u2026")
m.gap(0.05)
m.label(ROLE_USER).line("Now turn that into a tweet.", highlight=True)
m.composer(pin=False)

# the frame: everything above the composer is what gets re-read
FRAME_T = MY + 0.38
FRAME_B = MY + m.h - 0.56
s.rect(MX - 0.13, FRAME_T, 0.035, FRAME_B - FRAME_T, fill=ACCENT)
for _y in (FRAME_T, FRAME_B - 0.035):
    s.rect(MX - 0.13, _y, 0.16, 0.035, fill=ACCENT)
RX, RW = MX + MW + 0.24, 3.06
s.text(RX, 1.90, RW, 0.30, "The context window",
       size=SZ_DIAGRAM, color=ACCENT, bold=True, font=BODY_FONT)
s.text(RX, 2.28, RW, 1.20,
       "Everything inside the bracket is sent again, in full, on every turn. "
       "The model remembers nothing between turns - it only has this.",
       size=SZ_CAPTION, color=INK, font=BODY_FONT, spacing=1.25)

# the part a student thinks of as "the prompt" is one line of the whole
s.arrow(MX + MW + 0.02, FRAME_B - 0.44, 0.18, 0.16, direction="left")
s.text(RX, 3.62, RW, 0.56,
       "What you just typed is one line of it.",
       size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT, spacing=1.2)

s.text(RX, 4.26, RW, 0.72,
       "Fill the window and something gives: the reply is cut short, or the "
       "oldest turns stop being sent.",
       size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT, spacing=1.25)

s = d.content("Every turn re-sends the whole thread")
m = s.mock(0.40, 1.56, 4.90, 3.30)
m.user("Summarise this report in five points.")
with m.bubble():
    m.assistant("1. Revenue is flat.  2. Churn is up…")
m.user("Now turn that into a tweet.")
m.gap(0.04)
m.label("Sent with the second turn")
m.line("The report, both questions, the answer", highlight=True)
m.composer(pin=False)
s.beside(m,
         "A model holds nothing between requests. Each turn re-sends everything "
         "above it, so a long thread is counted again with every message you add.")

s = d.content("Token limits - set by the model you pick")
s.compare(1.56,
          "Input limit", ["How much you may send",
                          "Prompt, thread, files, all of it",
                          "Hit it and the request is refused"],
          "Output limit", ["How much it may write back",
                           "Far smaller than the input limit",
                           "Hit it and the answer is cut off"],
          h=2.30)
s.caption(0.34, 4.02, 9.32,
          "These are not settings you tune - they come with the model you chose, "
          "so picking a model is picking a budget. Two failures that look alike: "
          "a refused request means you sent too much, a truncated answer means "
          "you left it too little room.")

s = d.content("Counting them without guessing")
_png, _ratio = code_image("""
import tiktoken

enc = tiktoken.get_encoding("o200k_base")
assert enc.decode(enc.encode("hello world")) == "hello world"
""", ASSETS, "tiktoken")
_c = picture(s, _png, 0.40, 1.66, 5.60, ratio=_ratio)
s.text(0.40, _c + 0.12, 5.60, 0.26,
       "Encode to count. Decode to prove nothing was lost.",
       size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT)
_p = SimpleNamespace(x=0.40, y=1.66, w=5.60, h=_c - 1.66)
s.beside(_p,
         "The tokeniser is a library, not a mystery: the same rules the model "
         "uses, available to you. Count before you send, and a limit stops being "
         "something you discover by hitting it.")

d.two_columns(
    "When the window fills up",
    ["The message is refused outright",
     "Or the oldest turns quietly drop out",
     "Or they are compressed into a summary you cannot read"],
    ["Send the relevant extract, not the whole file",
     "Work in chunks, then combine the results",
     "Carry a summary forward instead of the transcript",
     "Start a fresh conversation for a fresh task"],
    left_heading="What happens",
    right_heading="What you do about it",
)

# ============================================================== WHY IT MATTERS
s = d.why("a token has three prices, not one")
y = bar_row(s, 1.56, "Cached input", "an opening seen before", 2.30,
            "a fraction of input")
y = bar_row(s, y, "Input", "your prompt and thread", 3.50, "the baseline")
y = bar_row(s, y, "Output", "every word it writes back", 6.50,
            "several times the input rate", accent=True)
s.caption(0.34, y + 0.14, 9.32,
          "The bars show the ordering, not a scale. Providers publish these three "
          "rates separately and revise them constantly - the ordering is what has "
          "held, and it is the part worth memorising.")

s = d.content("Caching pays for the part it recognises")
_y = 1.58
_y = prefix_bar(s, 0.34, _y, 9.32, "First time", "nothing to recognise", 0.0,
                fresh_text="all of it new - full price")
_y = prefix_bar(s, 0.34, _y, 9.32, "Ask again", "same opening, new question", 0.72)
_y = prefix_bar(s, 0.34, _y, 9.32, "Edit the top",
                "one word changed early", 0.12, broken=True,
                fresh_text="the match broke here - full price again")
s.caption(0.34, _y + 0.10, 9.32,
          "The match runs from the very first character and stops at the first "
          "difference, so a date or a reordered line near the top throws the whole "
          "discount away. Stable material first; the part that changes, last.")

s = d.content("Which of these prices do you actually pay?")
s.compare(1.56,
          "On a ChatGPT subscription",
          ["You never see a per-token bill",
           "You get an allowance per model instead",
           "Allowances refill on a rolling window, with a weekly cap on top",
           "Hit one and you wait for the reset, or switch model"],
          "Building on the API",
          ["You pay per token, in the three buckets",
           "Caching is a real lever you control",
           "Prompt order becomes an engineering decision",
           "We come back to this properly later in the course"],
          h=2.52)
s.caption(0.34, 4.24, 9.32,
          "Same tokens underneath, billed two different ways. On a subscription the "
          "cost of a long thread is not money, it is your allowance - which is why "
          "a runaway conversation quietly eats your week.")

s = d.content("What that changes about how you prompt")
s.compare(1.56,
          "Cheaper than it looks", ["A long brief, pasted once",
                                    "Worked examples of what you want",
                                    "Standing instructions you reuse"],
          "Dearer than it looks", ["A long essay written back",
                                   "\"Give me ten versions\"",
                                   "Re-generating instead of editing"],
          h=2.00)
s.caption(0.34, 3.72, 9.32,
          "Be generous with the context you give, specific about the length you ask "
          "for, and consistent about the order you give it in. The expensive half "
          "of a conversation is the half you did not write.")

# ---------------------------------------------------------------- handoff
d.handoff(
    [("Count a sentence", "paste it into a tokeniser"),
     ("Try another language", "same meaning, more tokens"),
     ("Read a price table", "three columns, not one")],
    lead="That is the principle. Now we go and count some tokens for real.",
)

path = d.save(out)
print(f"wrote {path}  ({len(d.prs.slides._sldIdLst)} slides)")
