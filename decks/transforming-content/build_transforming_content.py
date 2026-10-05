#!/usr/bin/env python3
"""Build the principles deck for Transforming Content for Your Audience (was ELI5)."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / ".claude" / "skills" / "principles-deck" / "scripts"))
from deckkit import Deck, title_case  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from deckkit import (ACCENT, ARROW, INK, SOFT_INK, RULE, NODE_BG, NODE_ACC,  # noqa: E402
                     BODY_FONT, HEAD_FONT, SZ_CAPTION, _rgb)

# Visual helpers for this deck (icons are Lucide, ISC licence, rendered to assets/).
ASSETS = Path(__file__).resolve().parent / "assets"
TONES = {"pink": (ACCENT, ACCENT, "white"), "soft": (NODE_ACC, ACCENT, "pink"),
         "grey": ("ECECF0", RULE, "ink")}


def mark(s):
    """Shape count now; pass it to group() to bundle everything drawn since."""
    return len(s.raw.shapes)


def group(s, first, name=None):
    """Group every shape drawn since `first`, so it moves as one unit."""
    shapes = s.raw.shapes
    g = shapes.add_group_shape(list(shapes)[first:])
    if name:
        g.name = name
    return g


def glyph(s, x, y, size, icon, colour="pink"):
    """A bare icon."""
    s.raw.shapes.add_picture(str(ASSETS / f"i-{icon}-{colour}.png"),
                             Inches(x), Inches(y), Inches(size), Inches(size))


def oval(s, x, y, d, fill, line=None):
    c = s.raw.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d),
                               Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = _rgb(fill)
    if line:
        c.line.color.rgb = _rgb(line)
        c.line.width = Pt(1.0)
    else:
        c.line.fill.background()
    c.shadow.inherit = False
    return c


def badge(s, x, y, d, icon, tone="pink"):
    """A round icon badge: pink (solid), soft (pink tint) or grey."""
    fill, line, colour = TONES[tone]
    oval(s, x, y, d, fill, line)
    g = d * 0.54
    glyph(s, x + (d - g) / 2, y + (d - g) / 2, g, icon, colour)

out = Path(sys.argv[1] if len(sys.argv) > 1 else
           Path(__file__).resolve().parent / "transforming-content.pptx")
d = Deck()

d.title(
    "Transforming Content for Your Audience",
    "Turning the same information into the form a different reader actually needs.",
)

# One source fans out to several readers. A bus bar keeps every arrow straight.
s = d.what("same facts, a new form for a new reader")
out_x, out_w, out_h, out_gap, top = 3.90, 5.76, 0.62, 0.12, 1.58
outputs = [
    ("mail", "Email summary", "for the team"),
    ("presentation", "Board slide outline", "for the directors"),
    ("smile", "Plain-Language explainer", "for a new starter"),
    ("languages", "Spanish translation", "for a partner office"),
]
mids = [top + i * (out_h + out_gap) + out_h / 2 for i in range(len(outputs))]
src_h = 1.50
src_y = (mids[0] + mids[-1]) / 2 - src_h / 2
first = mark(s)
s.rect(0.34, src_y, 2.50, src_h, fill=NODE_ACC, rounded=True, radius=0.08,
       line=ACCENT, line_w=1.0)
badge(s, 0.34 + 1.25 - 0.28, src_y + 0.16, 0.56, "notebook-pen", "pink")
s.text(0.34, src_y + 0.80, 2.50, 0.30, title_case("Meeting notes"), size=14,
       color=INK, bold=True, font=BODY_FONT, align="c")
s.text(0.34, src_y + 1.10, 2.50, 0.26, title_case("rough, long, internal"),
       size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT, align="c")
group(s, first, "Source")
bus_x = 3.30
s.rect(2.84, (mids[0] + mids[-1]) / 2 - 0.006, bus_x - 2.84, 0.012, fill=ARROW,
       rounded=True)
s.rect(bus_x, mids[0], 0.012, mids[-1] - mids[0], fill=ARROW)
for m in mids:
    s.arrow(bus_x, m - 0.06, out_x - bus_x - 0.08, 0.12)
for i, (icon, label, reader) in enumerate(outputs):
    y = top + i * (out_h + out_gap)
    first = mark(s)
    s.rect(out_x, y, out_w, out_h, fill=NODE_BG, rounded=True, radius=0.06,
           line=RULE, line_w=0.75)
    badge(s, out_x + 0.12, y + 0.10, out_h - 0.20, icon, "soft")
    s.text(out_x + 0.70, y, 3.00, out_h, title_case(label), size=13, color=INK,
           bold=True, font=BODY_FONT, anchor="m")
    s.text(out_x + 3.60, y, out_w - 3.76, out_h, title_case(reader),
           size=SZ_CAPTION, color=SOFT_INK, font=BODY_FONT, align="r", anchor="m")
    group(s, first, label)
s.caption(0.34, 4.66, 9.32,
          "The information stays the same. The shape changes for whoever reads it next.")

# Six dials, drawn as sliders: each has two ends and a setting.
s = d.how("six dials you can turn")
dials = [
    ("users", "Audience", "Your Team", "The Board", 0.82),
    ("gauge", "Reading level", "Age Five", "Expert", 0.06),
    ("ruler", "Length", "One Line", "Full Brief", 0.30),
    ("layout-template", "Format", "Email", "Slides", 0.70),
    ("languages", "Language", "English", "Any Language", 0.50),
    ("message-circle", "Tone", "Friendly", "Formal", 0.40),
]
cw, cg = 2.92, 0.28
for i, (icon, label, lo, hi, pos) in enumerate(dials):
    x = 0.34 + (i % 3) * (cw + cg)
    y = 1.56 + (i // 3) * 1.46
    first = mark(s)
    badge(s, x, y, 0.46, icon, "pink" if label == "Reading level" else "soft")
    s.text(x + 0.58, y, cw - 0.58, 0.46, title_case(label), size=14, color=INK,
           bold=True, font=BODY_FONT, anchor="m")
    ty = y + 0.78
    s.rect(x, ty, cw, 0.05, fill=RULE, rounded=True)
    s.rect(x, ty, cw * pos + 0.001, 0.05, fill=ACCENT, rounded=True)
    oval(s, x + cw * pos - 0.10, ty - 0.075, 0.20, "FFFFFF", ACCENT)
    s.text(x, ty + 0.16, cw / 2, 0.26, lo, size=SZ_CAPTION, color=SOFT_INK,
           font=BODY_FONT)
    s.text(x + cw / 2, ty + 0.16, cw / 2, 0.26, hi, size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT, align="r")
    group(s, first, f"Dial: {label}")
s.caption(0.34, 4.58, 9.32,
          "\"Explain it like I'm five\" is one setting: reading level turned all the way down.")

# The facts survive: the mock, with the three facts checked off beside it.
s = d.why("the facts survive, the shape changes")
m = s.mock(0.34, 1.56, 4.90, 3.44)
m.user("Turn these notes into a short email for the team.",
       "- q3 launch slips 2 wks (vendor)",
       "- Priya owns QA plan, due Fri",
       "- budget ok, +£4k contingency")
with m.bubble():
    m.label("AI Assistant")
    m.line("Subject: Q3 launch moves back two weeks")
    m.line("A vendor delay pushes the launch by two weeks.", highlight=True)
    m.line("Budget holds, with £4k in contingency.", highlight=True)
    m.line("Priya will share the QA plan by Friday.", highlight=True)
m.fit()
fx = 5.62
s.text(fx, 1.56, 4.04, 0.32, title_case("Facts that must survive"), size=14,
       color=INK, bold=True, font=BODY_FONT)
for i, (label, detail) in enumerate([("The date", "Two weeks later"),
                                     ("The owner", "Priya, by Friday"),
                                     ("The figure", "£4k contingency")]):
    y = 2.04 + i * 0.66
    first = mark(s)
    badge(s, fx, y, 0.44, "circle-check", "pink")
    s.text(fx + 0.58, y - 0.02, 3.40, 0.26, title_case(label), size=SZ_CAPTION,
           color=INK, bold=True, font=BODY_FONT)
    s.text(fx + 0.58, y + 0.22, 3.40, 0.26, detail, size=SZ_CAPTION,
           color=SOFT_INK, font=BODY_FONT)
    group(s, first, label)
s.caption(fx, 4.08, 4.04, "Check each one made it across.")

# The prompt pattern as a labelled form, not a block of text.
s = d.content("The prompt pattern")
fields = [
    ("users", "Audience", "{who will read it}"),
    ("target", "Purpose", "{what they should do or decide}"),
    ("layout-template", "Format", "{email, slide outline, FAQ}"),
    ("ruler", "Length and tone", "{under 150 words, plain}"),
    ("lock", "Keep exactly", "{numbers, names, dates, decisions}"),
    ("notebook-pen", "Content", "\"\"\"{paste the source here}\"\"\""),
]
s.text(0.34, 1.58, 9.32, 0.30, "Transform the content below.", size=SZ_CAPTION,
       color=INK, bold=True, font=BODY_FONT)
for i, (icon, label, value) in enumerate(fields):
    y = 1.98 + i * 0.44
    keep = label == "Keep exactly"
    first = mark(s)
    if keep:
        s.rect(0.34, y - 0.02, 9.32, 0.42, fill=NODE_ACC, rounded=True,
               radius=0.08, line=ACCENT, line_w=0.75)
    badge(s, 0.44, y + 0.03, 0.34, icon, "pink" if keep else "soft")
    s.text(0.92, y, 2.30, 0.38, title_case(label), size=SZ_CAPTION,
           color=ACCENT if keep else INK, bold=True, font=BODY_FONT, anchor="m")
    s.text(3.30, y, 6.30, 0.38, value, size=SZ_CAPTION, color=SOFT_INK,
           font=BODY_FONT, anchor="m")
    group(s, first, label)
s.caption(0.34, 4.72, 9.32,
          "Simpler can mean wrong. Ask it to list anything it left out or rounded off.")

d.handoff(
    [
        ("Paste notes", "a messy meeting Write-Up"),
        ("Ask for email", "audience, purpose, length"),
        ("Board version", "same notes, slide outline"),
        ("Check facts", "numbers, names, dates"),
    ],
    lead="Next, one set of meeting notes becomes two outputs in ChatGPT.",
)

path = d.save(out)
print(f"wrote {path}  ({len(d.prs.slides._sldIdLst)} slides)")
