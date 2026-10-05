#!/usr/bin/env python3
"""
Principles half of the NEW lecture "What is MCP?".

Build:
    uv run --with python-pptx python decks/what-is-mcp/build_what_is_mcp.py \
        decks/what-is-mcp/what-is-mcp.pptx
    uv run --with python-pptx python .claude/skills/principles-deck/scripts/check_deck.py \
        decks/what-is-mcp/what-is-mcp.pptx

Deliberately vendor-neutral: MCP is an open standard, not one lab's plug-in
system. No JSON, no transport names, no menu paths, no server directory - those
belong in the demo half. See NOTES.md.
"""
import sys
from pathlib import Path

from pptx.enum.shapes import MSO_CONNECTOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / ".claude" / "skills" / "principles-deck" / "scripts"))
from deckkit import (Deck, ACCENT, ARROW, INK, RULE, SOFT_INK,  # noqa: E402
                     BODY_FONT, SZ_DIAGRAM)

out = Path(sys.argv[1] if len(sys.argv) > 1
           else ROOT / "decks/what-is-mcp/what-is-mcp.pptx")
d = Deck()


def wire(slide, x1, y1, x2, y2, color=RULE, width=1.0):
    """A straight hairline between two points. deckkit only draws axis-aligned
    rules, and the whole point of the N x M slide is the diagonals."""
    c = slide.raw.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = __import__("deckkit")._rgb(color)
    c.line.width = Pt(width)
    c.shadow.inherit = False
    return c


# ---------------------------------------------------------------- open
d.title("What is MCP?",
        "How an assistant plugs into the software you already use, through one "
        "shared standard instead of a bespoke integration each time.")

# ============================================================== WHAT IS IT?
s = d.what("a common standard for reaching your other systems")
m = s.mock(0.40, 1.56, 5.10, 3.40)
m.user("Total last month's invoices by client.")
m.label("Calling a connected system")
m.line("Example Accounting - invoices, last month", highlight=True)
with m.bubble():
    m.assistant("Six invoices, £14,200 in total.",
                "Three clients account for most of it.")
m.composer(pin=False)
s.beside(m, "MCP - the Model Context Protocol - is an open standard for "
            "connecting an assistant to other software. The system at the other "
            "end runs a server, which lists the jobs it can do. Any assistant "
            "that speaks the standard can use any server.")

s = d.content("A skill is not a connection")
s.compare(1.56,
          "A skill", ["Instructions the assistant follows",
                      "Written in words, by you",
                      "Changes how it works",
                      "Travels with the assistant"],
          "MCP", ["A live link to another system",
                  "Spoken between programs",
                  "Changes what it can reach",
                  "Lives wherever that system runs"],
          h=2.40)
s.caption(0.34, 4.12, 9.32,
          "A \"connector\" or an \"app\" is usually an MCP server: the small "
          "program a system runs so assistants can reach it. A \"plugin\" is the "
          "bundle shipped around one - a server plus the skills for using it. The "
          "two stack rather than compete.")

s = d.content("The problem it was invented to solve")
PY_TOP, PY_H = 1.56, 2.60
for px, ptitle, accent in ((0.34, "Every pair, wired by hand", False),
                           (5.20, "One protocol in the middle", True)):
    s.rect(px, PY_TOP, 4.46, PY_H, fill="FFFFFF", rounded=True,
           line=ACCENT if accent else RULE, line_w=1.0 if accent else 0.75)
    s.text(px + 0.20, PY_TOP + 0.14, 4.06, 0.30, ptitle, size=14,
           color=ACCENT if accent else INK, bold=True, font=BODY_FONT)

ROWS = (2.06, 2.62, 3.18)
NH = 0.44
# left panel: three assistants, three systems, every pair wired
for i, y in enumerate(ROWS):
    s.node(0.56, y, 1.28, NH, f"Assistant {'ABC'[i]}")
    s.node(3.26, y, 1.28, NH, f"System {i + 1}")
for ya in ROWS:
    for yb in ROWS:
        wire(s, 1.84, ya + NH / 2, 3.26, yb + NH / 2)

# right panel: the same six, through one hub
HUB_TOP, HUB_BOT = ROWS[0], ROWS[-1] + NH
HUB_MID = (HUB_TOP + HUB_BOT) / 2
for i, y in enumerate(ROWS):
    s.node(5.42, y, 1.20, NH, f"Assistant {'ABC'[i]}")
    s.node(8.30, y, 1.20, NH, f"System {i + 1}")
    wire(s, 6.62, y + NH / 2, 7.10, HUB_MID, color=ARROW, width=1.25)
    wire(s, 7.96, HUB_MID, 8.30, y + NH / 2, color=ARROW, width=1.25)
s.node(7.10, HUB_TOP, 0.86, HUB_BOT - HUB_TOP, "MCP", accent=True)

s.caption(0.34, 4.32, 9.32,
          "Three assistants and three systems is nine integrations to build and "
          "keep working; four and ten is forty. Through one protocol each side "
          "learns the standard once, so the work adds up instead of "
          "multiplying - and no single company owns the plug.")

# ============================================================== HOW IT WORKS
s = d.how("tools, resources and prompts")
s.table(0.34, 1.56, 9.32, 2.20, [
    ["", "What it is", "For example"],
    ["Tools", "Actions the assistant can take", "Raise an invoice"],
    ["Resources", "Data it is allowed to read", "This month's invoices"],
    ["Prompts", "Ready-made instructions the server supplies",
     "Draft a chase-up email"],
], col_widths=[1.6, 4.4, 3.3])
s.caption(0.34, 3.96, 9.32,
          "One connection, three kinds of thing. Tools are the ones that change "
          "something in the world, so they are the ones a well-behaved "
          "assistant stops and asks you about. Reading is quiet; doing should "
          "not be.")

s = d.content("It finds out what it can do when it connects")
s.flow(0.34, 1.56, 9.32,
       [("Connect", "assistant to server"),
        ("Discover", "what the server offers"),
        ("Ask", "in plain language"),
        ("Call", "a tool, and read the result")],
       h=1.15, accent_last=True)
s.caption(0.34, 3.66, 9.32,
          "Nothing is hard-wired. The assistant finds out what a server can do "
          "at the moment it connects, which is why a system built last week "
          "works in an assistant built last year. That discovery step is the "
          "standard's real trick.")

s = d.content("What actually reaches the model")
bottom = s.layers(0.34, 1.56, 5.90,
                  [("The tools on offer", "listed by the server"),
                   ("Your message", "typed just now"),
                   ("The tool result", "sent back by the server", True)], h=0.66)
s.arrow(3.10, bottom + 0.04, 0.34, 0.22, direction="down")
s.node(0.34, bottom + 0.36, 5.90, 0.62,
       "The model reads all three the same way", accent=True)
s.caption(6.50, 1.60, 3.16,
          "A tool's description and a tool's result are text, written by "
          "whoever runs that server, landing in the same window as your own "
          "instructions. The model cannot tell them apart.")

s = d.content("Where a connection stops")
m = s.mock(5.10, 1.56, 4.56, 3.10, app="Example Accounting", tag="Connected")
m.label("It can")
m.line("List invoices you already have access to.")
m.line("Raise a draft invoice, once approved.")
m.gap(0.06)
m.divider()
m.label("It cannot")
m.line("Read your other conversations.")
m.line("Reach a system you never connected.")
m.line("Do more than your own account may do.")
m.fit()
s.beside(m, "A connection is narrow in both directions. The server sees the "
            "tool calls sent to it, not your whole conversation, and it acts "
            "with the permissions you granted - which are yours, not the "
            "assistant's.")
s.caption(0.34, 4.14, 9.32,
          "The boundary only protects you in one direction. What a server "
          "offers is the server's to change, so a tool can appear, vanish or "
          "quietly start doing something different, and nothing in your "
          "conversation announces it.")

# ============================================================== WHY IT MATTERS
s = d.why("built once, works everywhere")
s.table(0.34, 1.56, 9.32, 2.35, [
    ["", "A bespoke integration", "Through MCP"],
    ["Built for", "One assistant", "Any assistant that speaks it"],
    ["If you switch tools", "Built again from scratch", "Reconnect and carry on"],
    ["A capability you want", "Wait for the vendor", "Connect a server yourself"],
    ["Failure mode", "Nothing connects", "Too much connects"],
], col_widths=[2.1, 3.6, 3.6])
s.caption(0.34, 4.14, 9.32,
          "That last row is the change worth noticing. The old complaint was "
          "that your assistant could not reach your systems. The new problem is "
          "deciding which of them it should.")

s = d.content("Judging a server before you connect it")
s.steps(0.34, 1.56, 5.60, [
    "Who runs it - the service itself, or a stranger's copy?",
    "It acts as you, with your access, not the assistant's.",
    "Can it only read, or can it also write?",
    "Its words arrive as instructions the model may follow.",
])
s.caption(6.20, 1.60, 3.46,
          "A connected server is not a feature you switched on. It is a third "
          "party holding a set of your permissions, and everything it says "
          "arrives in the same window as everything you say.\n\n"
          "You will almost never write one of these. Your job is choosing which "
          "to trust: connect few, prefer the server the service runs itself, "
          "and allow writing only where you would allow a new colleague.")

d.two_columns(
    "When a connection earns its place",
    ["The work needs live data from a system you own",
     "You run the same lookup every week",
     "The assistant has to act, not only advise"],
    ["A single file upload would do",
     "It is a one-off you will never repeat",
     "You would not hand a new colleague the same access"],
    left_heading="Connect a server when",
    right_heading="Do not bother when",
)

# ---------------------------------------------------------------- handoff
d.handoff(
    [("Open the connections list", "and see what is already there"),
     ("Connect one server", "and read what it asks you for"),
     ("Ask something that needs it", "and watch the call come back")],
    lead="That is the principle. Now we plug one in and watch it work.",
)

path = d.save(out)
print(f"wrote {path}  ({len(d.prs.slides._sldIdLst)} slides)")
